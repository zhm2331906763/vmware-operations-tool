import os
import time
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import concurrent.futures
from src.db.database import Database
from src.core.vcenter_connector import VCenterConnector
from src.core.vm_cloner import VMCloner
from src.core.concurrent_config import ConcurrentConfig
from src.ui.theme import COLORS, FONTS
from src.ui.enhanced_widgets import EnhancedButton, ProgressDialog
from openpyxl import load_workbook


class ClonePage:
    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.db = Database()
        self._create_ui()

    def _create_ui(self):
        file_frame = ttk.Frame(self.parent)
        file_frame.pack(fill=tk.X, padx=10, pady=5)

        ttk.Label(file_frame, text="克隆配置:").pack(side=tk.LEFT)
        self.file_path_var = tk.StringVar()
        ttk.Entry(file_frame, textvariable=self.file_path_var, width=50).pack(side=tk.LEFT, padx=5)
        EnhancedButton(file_frame, "浏览", "📂", command=self._browse_file).pack(side=tk.LEFT)
        EnhancedButton(file_frame, "下载模板", "📥", command=self._download_template).pack(side=tk.LEFT, padx=5)

        btn_frame = ttk.Frame(self.parent)
        btn_frame.pack(fill=tk.X, padx=10, pady=5)

        self.execute_btn = EnhancedButton(btn_frame, "开始克隆", "▶️", command=self._execute_clone)
        self.execute_btn.pack(side=tk.LEFT)

        self.progress_frame = ttk.LabelFrame(self.parent, text="📊 克隆进度", padding=10, style='Card.TLabelframe')
        self.progress_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        self.progress_text = tk.Text(self.progress_frame, height=20, width=80, font=('Consolas', 9))
        self.progress_text.pack(fill=tk.BOTH, expand=True)

        vsb = ttk.Scrollbar(self.progress_frame, orient="vertical", command=self.progress_text.yview)
        self.progress_text.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        self.progress_text.tag_configure('success', foreground='#4CAF50')
        self.progress_text.tag_configure('error', foreground='#F44336')
        self.progress_text.tag_configure('info', foreground='#2196F3')
        self.progress_text.tag_configure('header', foreground='#1976D2', font=('Consolas', 9, 'bold'))

        status_frame = ttk.Frame(self.parent)
        status_frame.pack(fill=tk.X, padx=10, pady=(0, 5))

        self.status_label = ttk.Label(status_frame, text="✅ 就绪", font=FONTS['small'], anchor=tk.W)
        self.status_label.pack(fill=tk.X)

    def _browse_file(self):
        file_path = filedialog.askopenfilename(
            parent=self.parent,
            title="选择克隆配置Excel",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if file_path:
            self.file_path_var.set(file_path)

    def _download_template(self):
        template_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'templates')
        template_path = os.path.join(template_dir, 'vm_clone.xlsx')

        if os.path.exists(template_path):
            save_path = filedialog.asksaveasfilename(
                parent=self.parent,
                title="保存模板",
                defaultextension=".xlsx",
                filetypes=[("Excel文件", "*.xlsx")]
            )
            if save_path:
                import shutil
                shutil.copy(template_path, save_path)
                messagebox.showinfo("提示", f"模板已保存: {save_path}")
        else:
            messagebox.showwarning("提示", "模板文件不存在")

    def _log(self, message, tag='normal'):
        self.progress_text.insert(tk.END, message + '\n', tag)
        self.progress_text.see(tk.END)
        self.parent.update()

    def _power_on_with_retry(self, connector, vm_name, retries=3, base_delay=2):
        last_error = "开机失败"
        for attempt in range(1, retries + 1):
            try:
                vm = connector.get_vm_by_name(vm_name)
                if not vm:
                    last_error = "未找到已克隆的虚拟机"
                else:
                    success, msg = connector.power_on_vm(vm)
                    if success:
                        return True, msg
                    last_error = msg or last_error

                if attempt < retries:
                    time.sleep(base_delay * attempt)
            except Exception as e:
                last_error = str(e)
                if attempt < retries:
                    time.sleep(base_delay * attempt)
        return False, last_error

    def _clone_with_retry(self, cloner, config, retries=2, base_delay=3):
        last_error = "克隆失败"
        transient_markers = ("HostConnectFault", "Cannot connect to host", "connect to host")

        for attempt in range(1, retries + 1):
            success, msg = cloner.clone_vm(config)
            if success:
                return True, msg

            last_error = msg or last_error
            if attempt < retries and any(marker.lower() in str(last_error).lower() for marker in transient_markers):
                time.sleep(base_delay * attempt)
                continue
            break

        return False, last_error

    def _execute_clone(self):
        file_path = self.file_path_var.get()
        if not file_path:
            messagebox.showwarning("提示", "请选择克隆配置Excel文件")
            return

        try:
            configs = self._read_excel(file_path)
            if not configs:
                messagebox.showwarning("提示", "Excel文件中没有有效数据")
                return
        except Exception as e:
            messagebox.showerror("错误", f"读取Excel失败: {str(e)}")
            return

        self.execute_btn.config(state='disabled')
        self.progress_text.delete('1.0', tk.END)

        def run_clone():
            start_time = time.time()
            from collections import defaultdict
            vc_groups = defaultdict(list)
            for config in configs:
                vc_groups[config['vc_ip']].append(config)

            total = len(configs)
            success = 0
            failed = 0
            lock = threading.Lock()

            def process_single_vc(args):
                vc_ip, config_list = args
                vcenters = self.db.get_vcenters()
                vc_config = next((v for v in vcenters if v[1] == vc_ip), None)
                if not vc_config:
                    self.parent.after(0, lambda: self._log(f"[❌ 错误] 未找到VC配置: {vc_ip}", 'error'))
                    return len(config_list), len(config_list)

                self.parent.after(0, lambda: self._log(f"\n{'='*50}", 'header'))
                self.parent.after(0, lambda: self._log(f"🚀 开始处理 vCenter: {vc_ip}", 'header'))
                self.parent.after(0, lambda: self._log(f"{'='*50}", 'header'))

                local_success = 0
                local_failed = 0

                def clone_single(config):
                    self.parent.after(0, lambda: self._log(f"[{config['vm_name']}] ⏳ 正在克隆...", 'info'))
                    task_connector = VCenterConnector(vc_ip, vc_config[2], vc_config[3])
                    try:
                        success_conn, msg = task_connector.connect()
                        if not success_conn:
                            return config['vm_name'], False, f"连接VC失败: {msg}"

                        cloner = VMCloner(task_connector)
                        success_clone, msg_clone = self._clone_with_retry(cloner, config)

                        if success_clone:
                            time.sleep(2)
                            success_power_on, power_on_msg = self._power_on_with_retry(task_connector, config['vm_name'], retries=4, base_delay=3)
                            if success_power_on:
                                msg_clone += " ✅ (已自动开机)"
                            else:
                                msg_clone += f" ⚠️ (开机失败: {power_on_msg})"

                        return config['vm_name'], success_clone, msg_clone
                    finally:
                        task_connector.disconnect()

                max_workers = ConcurrentConfig.get_optimal_workers('clone_vm')
                max_workers = min(max_workers, len(config_list))

                with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = [executor.submit(clone_single, config) for config in config_list]
                    for future in concurrent.futures.as_completed(futures):
                        try:
                            vm_name, ok, msg = future.result()
                        except Exception as e:
                            vm_name = "未知虚拟机"
                            ok = False
                            msg = str(e)
                        if ok:
                            self.parent.after(0, lambda vm_name=vm_name, msg=msg: self._log(f"  ✅ {vm_name}: {msg}", 'success'))
                            local_success += 1
                        else:
                            self.parent.after(0, lambda vm_name=vm_name, msg=msg: self._log(f"  ❌ {vm_name}: {msg}", 'error'))
                            local_failed += 1

                self.parent.after(0, lambda: self._log(f"vCenter {vc_ip} 处理完成\n", 'header'))
                return local_success, local_failed

            max_workers = ConcurrentConfig.get_optimal_workers('vc_inspect')
            max_workers = min(max_workers, len(vc_groups))

            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = [executor.submit(process_single_vc, (vc_ip, configs)) for vc_ip, configs in vc_groups.items()]
                for future in concurrent.futures.as_completed(futures):
                    s, f = future.result()
                    success += s
                    failed += f

            elapsed = time.time() - start_time
            minutes = int(elapsed // 60)
            seconds = elapsed % 60
            time_str = f"{minutes}分{seconds:.1f}秒" if minutes > 0 else f"{seconds:.1f}秒"

            self.parent.after(0, lambda: self.execute_btn.config(state='normal'))
            self.parent.after(0, lambda: self.status_label.config(text=f"✅ 完成: 成功={success}, 失败={failed}, 耗时={time_str}"))
            self.parent.after(0, lambda: self._log(f"\n{'='*50}", 'header'))
            self.parent.after(0, lambda: self._log(f"🏁 全部完成: 成功={success}, 失败={failed}, 耗时={time_str}", 'header'))
            self.parent.after(0, lambda: self._log(f"{'='*50}", 'header'))

        threading.Thread(target=run_clone, daemon=True).start()

    def _read_excel(self, file_path):
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            return []

        configs = []
        for row in rows[1:]:
            if len(row) < 18:
                continue

            disk_sizes_str = str(row[16]).strip() if row[16] else ''
            disk_sizes = []
            if disk_sizes_str:
                for size_str in disk_sizes_str.split('\n'):
                    size_str = size_str.strip()
                    if size_str and size_str.isdigit():
                        disk_sizes.append(int(size_str))

            def safe_str(val, default=''):
                return str(val).strip() if val else default

            def safe_int(val):
                try:
                    return int(val) if val and str(val).strip().isdigit() else None
                except (ValueError, TypeError):
                    return None

            config = {
                'vc_ip': safe_str(row[0]),
                'host_name': safe_str(row[1]),
                'datastore_name': safe_str(row[2]),
                'folder_name': safe_str(row[3]),
                'template_name': safe_str(row[4]),
                'custom_spec_name': safe_str(row[5]) or None,
                'vm_name': safe_str(row[6]),
                'hostname': safe_str(row[7]),
                'cpu_num': safe_int(row[8]),
                'mem_num': safe_int(row[9]),
                'port_group': safe_str(row[10]),
                'nic_ip': safe_str(row[11]) or None,
                'nic_netmask': safe_str(row[12]) or '255.255.255.0',
                'nic_gateway': safe_str(row[13]) or None,
                'dns1': safe_str(row[14]),
                'dns2': safe_str(row[15]),
                'disk_sizes': disk_sizes,
                'remark': safe_str(row[17])
            }

            if config['vc_ip'] and config['vm_name'] and config['hostname']:
                configs.append(config)

        return configs
