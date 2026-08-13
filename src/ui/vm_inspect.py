import concurrent.futures
import threading
import time
import traceback

import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox

from src.core.concurrent_config import ConcurrentConfig
from src.core.vcenter_connector import VCenterConnector
from src.core.vm_collector import VMCollector
from src.core.vm_operations import VMNetworkAdapterManager
from src.core.vm_disk_manager import VMDiskManager
from src.core.vm_migrator import VMMigrator
from src.core.vm_operations import VMRenamer, VMRemarkUpdater
from src.core.vm_resize_manager import VMResizeManager
from src.db.database import Database
from src.core.excel_formatter import format_excel_file
from src.ui.enhanced_widgets import EnhancedButton, FilterableTable, ProgressDialog


class VMInspectPage:
    DISPLAY_COLUMNS = ('vCenter', '虚拟机名', '集群', '主机', '存储', '是否受NSX纳管')
    COLUMN_WIDTHS = {
        'vCenter': 120,
        '虚拟机名': 220,
        '集群': 140,
        '主机': 160,
        '存储': 260,
        '是否受NSX纳管': 120,
    }
    DB_COLUMN_INDICES = [1, 2, 3, 4, 17, 20]

    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.db = Database()
        self.vc_list = {}
        self.all_data = []
        self._create_ui()

    def _create_ui(self):
        toolbar = ttk.Frame(self.parent, style='Card.TFrame')
        toolbar.pack(fill=tk.X, padx=10, pady=5)

        vc_frame = ttk.Frame(toolbar)
        vc_frame.pack(side=tk.LEFT)
        ttk.Label(vc_frame, text='vCenter:').pack(side=tk.LEFT)
        self.vc_combo = ttk.Combobox(vc_frame, width=20, state='readonly')
        self.vc_combo.pack(side=tk.LEFT, padx=5)
        self.vc_combo.bind('<<ComboboxSelected>>', lambda e: self._on_vc_selected())

        btn_frame = ttk.Frame(toolbar)
        btn_frame.pack(side=tk.LEFT, padx=10)
        EnhancedButton(btn_frame, '开始巡检', '🔍', command=self._start_inspect).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, '一键巡检', '🚀', command=self._start_all_inspect).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, '导出Excel', '📄', command=self._export_excel).pack(side=tk.LEFT, padx=2)

        self.filterable_tree = FilterableTable(
            self.parent,
            columns=self.DISPLAY_COLUMNS,
            column_widths=self.COLUMN_WIDTHS,
            page_size=50,
            filter_columns=('vCenter', '虚拟机名', '集群', '主机', '存储', '是否受NSX纳管'),
            row_style_callback=self._get_row_tags,
        )
        self.filterable_tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        self.filterable_tree.bind_context_menu(self._show_context_menu)

        self._load_vc_list()
        self._load_all_data()
        self.parent.bind('<Visibility>', lambda e: self._refresh_vc_list())

    def _get_row_tags(self, row):
        if len(row) > 5:
            nsx_flag = str(row[5]).strip().lower()
            if '是' in nsx_flag or nsx_flag in {'yes', 'true', '1', 'y', 'nsx'}:
                return ('nsx_managed',)
        return ()

    def _show_context_menu(self, event):
        item_id = self.filterable_tree.tree.identify_row(event.y)
        if item_id:
            self.filterable_tree.tree.focus(item_id)
            if not self.filterable_tree.tree.selection():
                self.filterable_tree.tree.selection_set(item_id)

        selected = self.filterable_tree.get_selected()
        if not selected:
            return

        menu = tk.Menu(self.parent, tearoff=0)
        menu.add_command(label='虚拟机迁移', command=self._open_migrate_dialog)
        menu.add_command(label='添加网络适配器', command=self._open_add_nic_dialog)
        menu.add_command(label='修改备注', command=self._open_remark_dialog)
        menu.add_command(label='批量重命名', command=self._open_rename_dialog)
        menu.add_command(label='添加硬盘', command=self._open_add_disk_dialog)
        menu.add_command(label='CPU内存扩容', command=self._open_resize_dialog)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _copy_to_clipboard(self, text):
        value = str(text or '')
        self.parent.clipboard_clear()
        self.parent.clipboard_append(value)
        self.parent.update_idletasks()

    def _get_selected_vm_items(self):
        selected = self.filterable_tree.get_all_selected() or []
        items = []
        seen = set()
        for row in selected:
            if len(row) < 2:
                continue
            vc_ip = str(row[0]).strip()
            vm_name = str(row[1]).strip()
            if not vc_ip or not vm_name:
                continue
            key = (vc_ip, vm_name)
            if key in seen:
                continue
            seen.add(key)
            items.append({'vc_ip': vc_ip, 'vm_name': vm_name, 'row': row})
        return items

    def _group_vm_items(self, vm_items):
        grouped = {}
        for item in vm_items:
            grouped.setdefault(item['vc_ip'], []).append(item)
        return grouped

    def _connect_vcenter(self, vc_ip):
        user_pwd = self.vc_list.get(vc_ip)
        if not user_pwd:
            return None, f'未找到 {vc_ip} 的 vCenter 登录信息'
        connector = VCenterConnector(vc_ip, user_pwd[0], user_pwd[1])
        success, msg = connector.connect()
        if not success:
            connector.disconnect()
            return None, msg
        return connector, ''

    def _run_vm_batch_operation(self, title, vm_items, handler):
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        total = len(vm_items)
        progress = ProgressDialog(self.parent, title, determinate=True, maximum=total)
        progress.set_determinate(total, 0, f'准备执行 {title}...')

        def run():
            success_count = 0
            fail_count = 0
            grouped = self._group_vm_items(vm_items)

            try:
                processed = 0
                for vc_ip, items in grouped.items():
                    progress.update_status(f'正在连接 {vc_ip}...')
                    connector = None
                    try:
                        connector, err = self._connect_vcenter(vc_ip)
                        if not connector:
                            for item in items:
                                processed += 1
                                fail_count += 1
                                progress.update_progress(processed, total, f"{item['vm_name']}: 连接失败 {err}")
                            continue

                        for item in items:
                            processed += 1
                            progress.update_progress(processed, total, f'正在处理 {item["vm_name"]}...')
                            ok, msg = handler(connector, item['vm_name'])
                            if ok:
                                success_count += 1
                            else:
                                fail_count += 1
                            progress.update_status(f"{item['vm_name']}: {msg}")
                    finally:
                        if connector:
                            connector.disconnect()

                self.parent.after(
                    0,
                    lambda: messagebox.showinfo(
                        '完成',
                        f'{title}完成\n成功: {success_count}\n失败: {fail_count}',
                    ),
                )
                self.parent.after(0, lambda: self._load_all_data(preserve_view=True))
            except Exception as e:
                print(f'[VMInspect] {title} 失败: {e}')
                print(traceback.format_exc())
                self.parent.after(0, lambda: messagebox.showerror('错误', f'{title}失败: {e}'))
            finally:
                progress.close()

        threading.Thread(target=run, daemon=True).start()

    def _open_migrate_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'虚拟机迁移 ({len(vm_items)}台)')
        dialog.geometry('460x220')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').grid(row=0, column=0, columnspan=2, sticky='w', pady=(0, 8))
        ttk.Label(container, text='目标主机:').grid(row=1, column=0, sticky='w', pady=6)
        host_var = tk.StringVar()
        ttk.Entry(container, textvariable=host_var, width=34).grid(row=1, column=1, sticky='ew', pady=6)
        ttk.Label(container, text='目标存储:').grid(row=2, column=0, sticky='w', pady=6)
        datastore_var = tk.StringVar()
        ttk.Entry(container, textvariable=datastore_var, width=34).grid(row=2, column=1, sticky='ew', pady=6)
        ttk.Label(container, text='可只填一个，另一个留空表示只迁移计算或存储。', foreground=('#666666')).grid(
            row=3, column=0, columnspan=2, sticky='w', pady=(4, 0)
        )
        container.columnconfigure(1, weight=1)

        def on_confirm():
            target_host = host_var.get().strip() or None
            target_datastore = datastore_var.get().strip() or None
            if not target_host and not target_datastore:
                messagebox.showwarning('提示', '请至少填写目标主机或目标存储')
                return
            dialog.destroy()

            def handler(connector, vm_name):
                migrator = VMMigrator(connector)
                return migrator.migrate_vm(vm_name, target_host, target_datastore)

            self._run_vm_batch_operation('虚拟机迁移', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=14, sticky='e')
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _open_add_nic_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'添加网络适配器 ({len(vm_items)}台)')
        dialog.geometry('480x240')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').grid(row=0, column=0, columnspan=2, sticky='w')
        ttk.Label(container, text='端口组/网络名称:').grid(row=1, column=0, sticky='w', pady=6)
        portgroup_var = tk.StringVar()
        ttk.Entry(container, textvariable=portgroup_var, width=32).grid(row=1, column=1, sticky='ew', pady=6)
        ttk.Label(container, text='网卡类型:').grid(row=2, column=0, sticky='w', pady=6)
        adapter_var = tk.StringVar(value='vmxnet3')
        adapter_combo = ttk.Combobox(container, textvariable=adapter_var, state='readonly', values=['vmxnet3', 'e1000e', 'e1000', 'vmxnet2'])
        adapter_combo.grid(row=2, column=1, sticky='ew', pady=6)

        def on_confirm():
            portgroup_name = portgroup_var.get().strip()
            if not portgroup_name:
                messagebox.showwarning('提示', '请输入端口组或网络名称')
                return
            adapter_type = adapter_var.get().strip() or 'vmxnet3'
            dialog.destroy()

            def handler(connector, vm_name):
                manager = VMNetworkAdapterManager(connector)
                return manager.add_network_adapter(vm_name, portgroup_name, adapter_type)

            self._run_vm_batch_operation('添加网络适配器', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=14, sticky='e')
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _open_remark_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'修改备注 ({len(vm_items)}台)')
        dialog.geometry('460x220')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').pack(anchor='w')
        ttk.Label(container, text='新备注:').pack(anchor='w', pady=(12, 4))
        remark_text = tk.Text(container, height=5, width=48)
        remark_text.pack(fill=tk.BOTH, expand=True)

        def on_confirm():
            remark = remark_text.get('1.0', tk.END).strip()
            dialog.destroy()

            def handler(connector, vm_name):
                updater = VMRemarkUpdater(connector)
                return updater.update_remark(vm_name, remark)

            self._run_vm_batch_operation('修改备注', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.pack(fill=tk.X, pady=10)
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _open_rename_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'批量重命名 ({len(vm_items)}台)')
        dialog.geometry('520x280')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)

        rule_var = tk.StringVar(value='replace')
        content_var = tk.StringVar()

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').grid(row=0, column=0, columnspan=2, sticky='w')
        rule_frame = ttk.LabelFrame(container, text='命名方式', padding=10)
        rule_frame.grid(row=1, column=0, columnspan=2, sticky='ew', pady=10)
        ttk.Radiobutton(rule_frame, text='前缀', variable=rule_var, value='prefix').pack(anchor='w')
        ttk.Radiobutton(rule_frame, text='后缀', variable=rule_var, value='suffix').pack(anchor='w')
        ttk.Radiobutton(rule_frame, text='替换为指定名称', variable=rule_var, value='replace').pack(anchor='w')

        ttk.Label(container, text='内容:').grid(row=2, column=0, sticky='w', pady=6)
        ttk.Entry(container, textvariable=content_var, width=36).grid(row=2, column=1, sticky='ew', pady=6)
        container.columnconfigure(1, weight=1)

        def on_confirm():
            content = content_var.get().strip()
            if not content:
                messagebox.showwarning('提示', '请输入命名内容')
                return
            rule = rule_var.get()
            dialog.destroy()

            def handler(connector, vm_name):
                renamer = VMRenamer(connector)
                if rule == 'prefix':
                    new_name = f'{content}{vm_name}'
                elif rule == 'suffix':
                    new_name = f'{vm_name}{content}'
                else:
                    new_name = content
                return renamer.rename_vm(vm_name, new_name)

            self._run_vm_batch_operation('批量重命名', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=14, sticky='e')
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _open_add_disk_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'添加硬盘 ({len(vm_items)}台)')
        dialog.geometry('520x260')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').grid(row=0, column=0, columnspan=2, sticky='w')
        ttk.Label(container, text='硬盘大小(GB):').grid(row=1, column=0, sticky='w', pady=6)
        size_text = tk.Text(container, height=5, width=38)
        size_text.grid(row=1, column=1, sticky='ew', pady=6)
        ttk.Label(container, text='多个大小可换行或用逗号分隔，例如 100,200', foreground=('#666666')).grid(
            row=2, column=0, columnspan=2, sticky='w'
        )
        ttk.Label(container, text='目标存储(可选):').grid(row=3, column=0, sticky='w', pady=6)
        datastore_var = tk.StringVar()
        ttk.Entry(container, textvariable=datastore_var, width=36).grid(row=3, column=1, sticky='ew', pady=6)

        def on_confirm():
            raw_sizes = size_text.get('1.0', tk.END).replace(',', '\n').splitlines()
            disk_sizes = []
            for text in raw_sizes:
                text = text.strip()
                if not text:
                    continue
                try:
                    disk_sizes.append(int(float(text)))
                except Exception:
                    continue
            if not disk_sizes:
                messagebox.showwarning('提示', '请输入有效的硬盘大小')
                return
            datastore = datastore_var.get().strip() or None
            dialog.destroy()

            def handler(connector, vm_name):
                manager = VMDiskManager(connector)
                results = manager.add_multiple_disks(vm_name, disk_sizes, datastore)
                ok = all(r[1] for r in results)
                msg = '; '.join([f'{r[0]}GB:{r[2]}' for r in results])
                return ok, msg

            self._run_vm_batch_operation('添加硬盘', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=14, sticky='e')
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _open_resize_dialog(self):
        vm_items = self._get_selected_vm_items()
        if not vm_items:
            messagebox.showwarning('提示', '请先选择虚拟机')
            return

        dialog = tk.Toplevel(self.parent)
        dialog.title(f'CPU内存扩容 ({len(vm_items)}台)')
        dialog.geometry('460x230')
        dialog.transient(self.parent)
        dialog.grab_set()

        container = ttk.Frame(dialog, padding=16)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)

        ttk.Label(container, text=f'已选择 {len(vm_items)} 台虚拟机').grid(row=0, column=0, columnspan=2, sticky='w')
        ttk.Label(container, text='CPU数量(可空):').grid(row=1, column=0, sticky='w', pady=6)
        cpu_var = tk.StringVar()
        ttk.Entry(container, textvariable=cpu_var, width=28).grid(row=1, column=1, sticky='ew', pady=6)
        ttk.Label(container, text='内存(GB，可空):').grid(row=2, column=0, sticky='w', pady=6)
        memory_var = tk.StringVar()
        ttk.Entry(container, textvariable=memory_var, width=28).grid(row=2, column=1, sticky='ew', pady=6)

        def parse_positive_int(value):
            text = str(value).strip()
            if not text:
                return None
            try:
                number = int(float(text))
                return number if number > 0 else None
            except Exception:
                return None

        def on_confirm():
            cpu_count = parse_positive_int(cpu_var.get())
            memory_gb = parse_positive_int(memory_var.get())
            if cpu_count is None and memory_gb is None:
                messagebox.showwarning('提示', 'CPU 和内存不能同时为空')
                return
            dialog.destroy()

            def handler(connector, vm_name):
                manager = VMResizeManager(connector)
                return manager.resize_vm_resources(vm_name, cpu_count=cpu_count, memory_gb=memory_gb)

            self._run_vm_batch_operation('CPU内存扩容', vm_items, handler)

        btn_frame = ttk.Frame(container)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=14, sticky='e')
        ttk.Button(btn_frame, text='取消', command=dialog.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='确定', command=on_confirm).pack(side=tk.RIGHT, padx=4)

    def _load_vc_list(self):
        vcenters = self.db.get_vcenters()
        self.vc_list = {vc[1]: (vc[2], vc[3]) for vc in vcenters}
        vc_options = ['全部'] + list(self.vc_list.keys())
        self.vc_combo['values'] = vc_options
        if vc_options:
            self.vc_combo.current(0)

    def _refresh_vc_list(self):
        current = self.vc_combo.get()
        vcenters = self.db.get_vcenters()
        self.vc_list = {vc[1]: (vc[2], vc[3]) for vc in vcenters}
        vc_options = ['全部'] + list(self.vc_list.keys())
        self.vc_combo['values'] = vc_options
        if current in vc_options:
            self.vc_combo.set(current)
        else:
            self.vc_combo.current(0)
        self._load_all_data()

    def _on_vc_selected(self):
        self._load_all_data()

    def _load_all_data(self, preserve_view=False):
        selected_vc = self.vc_combo.get()
        if selected_vc == '全部':
            data, _ = self.db.get_vm_inspect(None, '', 1, 100000)
        else:
            data, _ = self.db.get_vm_inspect(selected_vc, '', 1, 100000)

        filtered_data = []
        for row in data:
            new_row = tuple(row[i] for i in self.DB_COLUMN_INDICES)
            new_row = tuple('' if v is None else v for v in new_row)
            filtered_data.append(new_row)

        self.all_data = filtered_data
        self.filterable_tree.set_data(filtered_data, preserve_view=preserve_view)

    def _start_inspect(self):
        vc_ip = self.vc_combo.get()
        if not vc_ip or vc_ip == '全部':
            messagebox.showwarning('提示', '请先选择一个 vCenter')
            return

        user, pwd = self.vc_list[vc_ip]
        progress = ProgressDialog(self.parent, '虚拟机巡检中')

        def run_inspect():
            connector = None
            try:
                start_time = time.time()
                progress.update_status(f'正在连接 {vc_ip}...')

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()
                if not success:
                    progress.close()
                    self.parent.after(0, lambda: messagebox.showerror('错误', f'连接失败: {msg}'))
                    return

                progress.update_status('正在获取虚拟机数据...')
                collector = VMCollector(connector)

                def on_progress(current, total, message=None):
                    progress.update_progress(current, total, message or f'正在解析虚拟机数据... {current}/{total}')

                vms = collector.collect_all_vms(progress_callback=on_progress)

                progress.update_status('正在保存巡检结果...')
                self.db.save_vm_inspect(vms, vc_ip=vc_ip)

                elapsed = time.time() - start_time
                progress.close()
                self.parent.after(
                    0,
                    lambda: messagebox.showinfo('完成', f'巡检完成，共采集 {len(vms)} 台虚拟机\n耗时: {elapsed:.1f} 秒')
                )
                self.parent.after(0, self._load_all_data)
            except Exception as e:
                print(f'[VMInspect] 巡检失败 vc={vc_ip}: {e}')
                print(traceback.format_exc())
                progress.close()
                self.parent.after(0, lambda: messagebox.showerror('错误', f'巡检失败: {e}'))
            finally:
                if connector:
                    connector.disconnect()

        threading.Thread(target=run_inspect, daemon=True).start()

    def _start_all_inspect(self):
        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning('提示', '没有可巡检的 vCenter')
            return

        self.db.clear_vm_inspect()
        self.filterable_tree.set_data([])

        total_vcenters = len(vcenters)
        progress = ProgressDialog(self.parent, '一键巡检中', determinate=True, maximum=total_vcenters)
        progress.set_determinate(total_vcenters, 0, f'正在巡检 VC 0/{total_vcenters}')
        start_time = time.time()
        total_vms = [0]
        results = []

        def inspect_single_vc(vc):
            vc_ip, user, pwd = vc[1], vc[2], vc[3]
            connector = None
            try:
                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()
                if not success:
                    return {'vc_ip': vc_ip, 'success': False, 'error': msg, 'vms': 0}

                collector = VMCollector(connector)
                vms = collector.collect_all_vms()
                self.db.save_vm_inspect(vms, vc_ip=vc_ip)
                return {'vc_ip': vc_ip, 'success': True, 'vms': len(vms), 'error': None}
            except Exception as e:
                tb = traceback.format_exc()
                print(f'[VMInspect] VC 巡检失败 vc={vc_ip}: {e}')
                print(tb)
                return {'vc_ip': vc_ip, 'success': False, 'error': str(e), 'vms': 0, 'traceback': tb}
            finally:
                if connector:
                    connector.disconnect()

        def run_inspect():
            workers = min(ConcurrentConfig.get_optimal_workers('vc_inspect'), total_vcenters)
            with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
                futures = {executor.submit(inspect_single_vc, vc): vc for vc in vcenters}
                completed = 0
                for future in concurrent.futures.as_completed(futures):
                    try:
                        result = future.result()
                    except Exception as e:
                        tb = traceback.format_exc()
                        print(f'[VMInspect] future 异常: {e}')
                        print(tb)
                        vc = futures.get(future)
                        vc_ip = vc[1] if vc and len(vc) > 1 else 'unknown'
                        result = {'vc_ip': vc_ip, 'success': False, 'error': str(e), 'vms': 0, 'traceback': tb}
                    results.append(result)
                    total_vms[0] += result['vms']
                    completed += 1
                    progress.update_progress(completed, total_vcenters, f'正在巡检 VC {completed}/{total_vcenters}')

            elapsed = int(time.time() - start_time)
            minutes = elapsed // 60
            seconds = elapsed % 60
            time_str = f'{minutes}m{seconds}s' if minutes > 0 else f'{seconds}s'
            failed = sum(1 for r in results if not r.get('success'))

            progress.update_progress(total_vcenters, total_vcenters, '巡检完成')
            msg = (
                f'一键巡检完成\n'
                f'总虚拟机数: {total_vms[0]}\n'
                f'耗时: {time_str}'
            )
            if failed > 0:
                msg += f'\n失败的 vCenter 数: {failed}'
            self.parent.after(50, progress.close)
            self.parent.after(400, lambda: messagebox.showinfo('完成', msg))
            self.parent.after(400, self._refresh_vc_list)

        threading.Thread(target=run_inspect, daemon=True).start()

    def _export_excel(self):
        from tkinter import filedialog

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title='保存 Excel',
            defaultextension='.xlsx',
            filetypes=[('Excel 文件', '*.xlsx')]
        )
        if not file_path:
            return

        try:
            selected_vc = self.vc_combo.get()
            if selected_vc == '全部':
                data, _ = self.db.get_vm_inspect(None, '', 1, 100000)
            else:
                data, _ = self.db.get_vm_inspect(selected_vc, '', 1, 100000)

            if not data:
                messagebox.showwarning('提示', '没有数据')
                return

            df = pd.DataFrame(data, columns=[
                'ID', 'vCenter', '虚拟机名', '集群', '主机', '电源状态', 'IP地址',
                '内存(GB)', 'CPU数量', 'UUID', '备注', '系统版本', '客户机系统',
                'CPU热添加是否开启', '内存热添加是否开启', 'VMware Tools状态', '运行天数', '存储', '磁盘', '端口组', '是否受NSX纳管', '已用空间(GB)',
                '文件夹', '巡检时间'
            ])
            region_map = self.db.get_vcenter_region_map()
            df.insert(0, '区域', df['vCenter'].map(lambda x: region_map.get(str(x), '') if pd.notna(x) else ''))
            df.to_excel(file_path, index=False)
            from openpyxl import load_workbook
            from openpyxl.styles import Font
            wb = load_workbook(file_path)
            ws = wb.active
            header_map = {ws.cell(row=1, column=i).value: i for i in range(1, ws.max_column + 1)}
            red_font = Font(color='9C0006')
            for col_name in ('CPU热添加是否开启', '内存热添加是否开启'):
                col_idx = header_map.get(col_name)
                if not col_idx:
                    continue
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if str(cell.value).strip() == '否':
                        cell.font = red_font
            format_excel_file(
                file_path,
                highlight_column='是否受NSX纳管',
                highlight_values={'是', 'yes', 'true', '1', 'y', 'nsx'},
            )
            wb.close()
            messagebox.showinfo('提示', f'导出成功: {file_path}')
        except Exception as e:
            messagebox.showerror('错误', f'导出失败: {e}')
