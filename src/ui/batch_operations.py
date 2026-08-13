import os
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import concurrent.futures
from collections import defaultdict
from src.db.database import Database
from src.core.vcenter_connector import VCenterConnector
from src.core.iscsi_manager import ISCSIManager
from src.core.vm_migrator import VMMigrator
from src.core.vm_operations import VMRenamer, VMRemarkUpdater
from src.core.dvswitch_manager import DVSwitchManager
from src.core.concurrent_config import ConcurrentConfig
from src.ui.theme import COLORS, FONTS
from src.core.esxi_host_manager import ESXiHostManager
from src.core.vm_disk_manager import VMDiskManager
from src.core.vm_resize_manager import VMResizeManager
from src.ui.theme import COLORS, FONTS
from src.ui.enhanced_widgets import EnhancedButton, ProgressDialog
from openpyxl import load_workbook


class BatchOperationsPage:
    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.db = Database()
        self._create_ui()

    def set_operation(self, operation_key):
        if not hasattr(self, "operation_var"):
            return
        if operation_key:
            self.operation_var.set(operation_key)
            self._on_operation_change()

    def _create_ui(self):
        op_frame = ttk.LabelFrame(self.parent, text="批量操作类型", padding=10, style='Card.TLabelframe')
        op_frame.pack(fill=tk.X, padx=10, pady=5)

        self.operation_var = tk.StringVar(value="iscsi")

        ttk.Radiobutton(op_frame, text="🖴 添加iSCSI", variable=self.operation_var,
                       value="iscsi", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🔄 虚拟机迁移", variable=self.operation_var,
                       value="migrate", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="📝 修改备注", variable=self.operation_var,
                       value="remark", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="✏️ 批量重命名", variable=self.operation_var,
                       value="rename", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🔀 创建分布式交换机", variable=self.operation_var,
                       value="dvswitch", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🔐 ESXi证书续订", variable=self.operation_var,
                       value="cert", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🔌 开启SSH", variable=self.operation_var,
                       value="ssh", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🕐 配置NTP", variable=self.operation_var,
                       value="ntp", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="💾 添加硬盘", variable=self.operation_var,
                       value="add_disk", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(op_frame, text="🧩 CPU内存扩容", variable=self.operation_var,
                       value="resize", command=self._on_operation_change).pack(side=tk.LEFT, padx=10)

        file_frame = ttk.Frame(self.parent)
        file_frame.pack(fill=tk.X, padx=10, pady=5)

        ttk.Label(file_frame, text="导入文件:").pack(side=tk.LEFT)
        self.file_path_var = tk.StringVar()
        ttk.Entry(file_frame, textvariable=self.file_path_var, width=50).pack(side=tk.LEFT, padx=5)
        EnhancedButton(file_frame, "浏览", "📂", command=self._browse_file).pack(side=tk.LEFT)
        EnhancedButton(file_frame, "下载模板", "📥", command=self._download_template).pack(side=tk.LEFT, padx=5)

        btn_frame = ttk.Frame(self.parent)
        btn_frame.pack(fill=tk.X, padx=10, pady=5)

        self.execute_btn = EnhancedButton(btn_frame, "执行", "▶️", command=self._execute_operation)
        self.execute_btn.pack(side=tk.LEFT)

        result_frame = ttk.LabelFrame(self.parent, text="执行结果", padding=10, style='Card.TLabelframe')
        result_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        columns = ('序号', '状态', '详情')
        self.tree = ttk.Treeview(result_frame, columns=columns, show='headings')

        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=100)

        vsb = ttk.Scrollbar(result_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.tree.tag_configure('success', background='#E8F5E9')
        self.tree.tag_configure('failure', background='#FFEBEE')
        self.tree.tag_configure('info', background='#E3F2FD')

        status_frame = ttk.Frame(self.parent)
        status_frame.pack(fill=tk.X, padx=10, pady=(0, 5))

        self.status_label = ttk.Label(status_frame, text="✅ 就绪", font=FONTS['small'], anchor=tk.W)
        self.status_label.pack(fill=tk.X)

        self._on_operation_change()

    def _on_operation_change(self):
        op = self.operation_var.get()
        templates = {
            'iscsi': 'iscsi.xlsx',
            'migrate': 'vmotion.xlsx',
            'remark': 'modify_remark.xlsx',
            'rename': 'rename.xlsx',
            'dvswitch': 'distributed_switch.xlsx',
            'cert': 'esxi_cert.xlsx',
            'ssh': 'esxi_ssh.xlsx',
            'ntp': 'esxi_ntp.xlsx',
            'add_disk': 'vm_add_disk.xlsx',
            'resize': 'vm_resize.xlsx'
        }
        self.current_template = templates.get(op, '')

    def _browse_file(self):
        file_path = filedialog.askopenfilename(
            parent=self.parent,
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if file_path:
            self.file_path_var.set(file_path)

    def _download_template(self):
        template_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'templates')
        template_path = os.path.join(template_dir, self.current_template)

        if os.path.exists(template_path):
            save_path = filedialog.asksaveasfilename(
                parent=self.parent,
                title="保存模板",
                defaultextension=".xlsx",
                filetypes=[("Excel文件", "*.xlsx")]
            )
            if save_path:
                import shutil
                shutil.copy(template_path, save_path)
                messagebox.showinfo("提示", f"模板已保存: {save_path}")
        else:
            messagebox.showwarning("提示", "模板文件不存在")

    def _execute_operation(self):
        file_path = self.file_path_var.get()
        if not file_path:
            messagebox.showwarning("提示", "请选择Excel文件")
            return

        op = self.operation_var.get()

        try:
            if op == "dvswitch":
                data = self._read_dvswitch_excel(file_path)
                configs = [data] if data['switches'] or data['portgroups'] else []
            else:
                configs = self._read_excel(file_path)

            if not configs:
                messagebox.showwarning("提示", "Excel文件中没有有效数据")
                return
        except Exception as e:
            messagebox.showerror("错误", f"读取Excel失败: {str(e)}")
            return

        vc_groups = defaultdict(list)

        if op == "dvswitch":
            for config in configs:
                data = config
                for sw in data.get('switches', []):
                    vc_groups[sw['vc_ip']].append(data)
                if not data.get('switches', []):
                    for pg in data.get('portgroups', []):
                        vc_groups[pg['vc_ip']].append(data)
        else:
            for config in configs:
                vc_groups[config['vc_ip']].append(config)

        self.execute_btn.config(state='disabled')
        self.status_label.config(text="⏳ 正在执行...")

        for item in self.tree.get_children():
            self.tree.delete(item)

        def run_operation():
            total = len(configs)
            success = 0
            failed = 0
            lock = threading.Lock()

            def process_single_vc(args):
                vc_ip, config_list = args
                vcenters = self.db.get_vcenters()
                vc_config = next((v for v in vcenters if v[1] == vc_ip), None)
                if not vc_config:
                    for cfg in config_list:
                        self._add_result(len(self.tree.get_children()) + 1, "失败", f"VC {vc_ip} 未找到")
                    return len(config_list), len(config_list)

                connector = VCenterConnector(vc_ip, vc_config[2], vc_config[3])
                success_conn, msg = connector.connect()
                if not success_conn:
                    for cfg in config_list:
                        self._add_result(len(self.tree.get_children()) + 1, "失败", f"VC连接失败: {msg}")
                    return 0, len(config_list)

                local_success = 0
                local_failed = 0

                if op == "iscsi":
                    local_success, local_failed = self._execute_iscsi(connector, config_list)
                elif op == "migrate":
                    local_success, local_failed = self._execute_migrate(connector, config_list)
                elif op == "remark":
                    local_success, local_failed = self._execute_remark(connector, config_list)
                elif op == "rename":
                    local_success, local_failed = self._execute_rename(connector, config_list)
                elif op == "dvswitch":
                    local_success, local_failed = self._execute_dvswitch(connector, config_list)
                elif op == "cert":
                    local_success, local_failed = self._execute_cert(connector, config_list)
                elif op == "ssh":
                    local_success, local_failed = self._execute_ssh(connector, config_list)
                elif op == "ntp":
                    local_success, local_failed = self._execute_ntp(connector, config_list)
                elif op == "add_disk":
                    local_success, local_failed = self._execute_add_disk(connector, config_list)
                elif op == "resize":
                    local_success, local_failed = self._execute_resize(connector, config_list)

                connector.disconnect()
                return local_success, local_failed

            max_workers = ConcurrentConfig.get_optimal_workers('vc_inspect')
            max_workers = min(max_workers, len(vc_groups))

            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {executor.submit(process_single_vc, (vc_ip, configs)): vc_ip
                          for vc_ip, configs in vc_groups.items()}

                for future in concurrent.futures.as_completed(futures):
                    s, f = future.result()
                    success += s
                    failed += f

            self.parent.after(0, lambda: self.execute_btn.config(state='normal'))
            self.parent.after(0, lambda: self.status_label.config(text=f"✅ 完成: 成功={success}, 失败={failed}"))

        threading.Thread(target=run_operation, daemon=True).start()

    def _read_excel(self, file_path):
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            return []

        configs = []
        op = self.operation_var.get()

        if op == "iscsi":
            for row in rows[1:]:
                if len(row) >= 3 and row[0] and row[1] and row[2]:
                    ip_str = str(row[2]).strip()
                    target_ips = [ip.strip() for ip in ip_str.split(',') if ip.strip()]
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'esxi_host': str(row[1]).strip(),
                        'target_ips': target_ips
                    })

        elif op == "migrate":
            for row in rows[1:]:
                if len(row) >= 4 and row[0] and row[1]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'vm_name': str(row[1]).strip(),
                        'target_host': str(row[2]).strip() if row[2] else None,
                        'target_datastore': str(row[3]).strip() if row[3] else None
                    })

        elif op == "remark":
            for row in rows[1:]:
                if len(row) >= 3 and row[0] and row[1]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'vm_name': str(row[1]).strip(),
                        'remark': str(row[2]).strip() if row[2] else ''
                    })

        elif op == "rename":
            for row in rows[1:]:
                if len(row) >= 3 and row[0] and row[1]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'vm_name': str(row[1]).strip(),
                        'new_name': str(row[2]).strip()
                    })

        elif op == "cert":
            for row in rows[1:]:
                if len(row) >= 2 and row[0] and row[1]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'esxi_host': str(row[1]).strip()
                    })

        elif op == "ssh":
            for row in rows[1:]:
                if len(row) >= 2 and row[0] and row[1]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'esxi_host': str(row[1]).strip()
                    })

        elif op == "ntp":
            for row in rows[1:]:
                if len(row) >= 3 and row[0] and row[1] and row[2]:
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'esxi_host': str(row[1]).strip(),
                        'ntp_servers': str(row[2]).strip()
                    })

        elif op == "add_disk":
            for row in rows[1:]:
                if len(row) >= 4 and row[0] and row[1] and row[2] and row[3]:
                    vm_name = str(row[1]).strip()
                    disk_sizes_str = str(row[2]).strip()
                    datastore = str(row[3]).strip()
                    disk_sizes = [s.strip() for s in disk_sizes_str.split('\n') if s.strip()]
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'vm_name': vm_name,
                        'disk_sizes': disk_sizes,
                        'datastore': datastore
                    })
        elif op == "resize":
            for row in rows[1:]:
                if len(row) >= 2 and row[0] and row[1]:
                    cpu_count = self._parse_optional_positive_int(row[2] if len(row) > 2 else None)
                    memory_gb = self._parse_optional_positive_int(row[3] if len(row) > 3 else None)
                    if cpu_count is None and memory_gb is None:
                        continue
                    configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'vm_name': str(row[1]).strip(),
                        'cpu_count': cpu_count,
                        'memory_gb': memory_gb
                    })

        return configs

    def _read_dvswitch_excel(self, file_path):
        wb = load_workbook(file_path, read_only=True, data_only=True)

        switch_configs = []
        portgroup_configs = []

        if '分布式交换机' in wb.sheetnames:
            ws_switch = wb['分布式交换机']
            rows_switch = list(ws_switch.iter_rows(values_only=True))
            for row in rows_switch[1:]:
                if len(row) >= 3 and row[0] and row[1] and row[2]:
                    switch_configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'datacenter': str(row[1]).strip(),
                        'switch_name': str(row[2]).strip()
                    })

        if '分布式端口组' in wb.sheetnames:
            ws_pg = wb['分布式端口组']
            rows_pg = list(ws_pg.iter_rows(values_only=True))
            for row in rows_pg[1:]:
                if len(row) >= 4 and row[0] and row[1] and row[2]:
                    try:
                        vlan_id = int(row[3]) if row[3] and str(row[3]).strip().isdigit() else 0
                    except (ValueError, TypeError):
                        vlan_id = 0
                    portgroup_configs.append({
                        'vc_ip': str(row[0]).strip(),
                        'switch_name': str(row[1]).strip(),
                        'portgroup_name': str(row[2]).strip(),
                        'vlan_id': vlan_id
                    })

        wb.close()

        return {
            'switches': switch_configs,
            'portgroups': portgroup_configs
        }

    def _execute_iscsi(self, connector, config_list):
        manager = ISCSIManager(connector)
        success_count = 0
        fail_count = 0

        def iscsi_single(config):
            success, msg = manager.add_iscsi_target(
                esxi_host_name=config['esxi_host'],
                target_ips=config['target_ips'],
                target_iqn=config.get('target_iqn')
            )
            return (config['esxi_host'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(iscsi_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                host, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{host}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_migrate(self, connector, config_list):
        migrator = VMMigrator(connector)
        success_count = 0
        fail_count = 0

        def migrate_single(config):
            success, msg = migrator.migrate_vm(
                config['vm_name'],
                config.get('target_host'),
                config.get('target_datastore')
            )
            return (config['vm_name'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(migrate_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                vm_name, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{vm_name}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_remark(self, connector, config_list):
        updater = VMRemarkUpdater(connector)
        success_count = 0
        fail_count = 0

        def remark_single(config):
            success, msg = updater.update_remark(config['vm_name'], config['remark'])
            return (config['vm_name'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(remark_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                vm_name, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"主机名 {vm_name}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_rename(self, connector, config_list):
        renamer = VMRenamer(connector)
        success_count = 0
        fail_count = 0

        def rename_single(config):
            success, msg = renamer.rename_vm(config['vm_name'], config['new_name'])
            return (config['vm_name'], config['new_name'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(rename_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                vm_name, new_name, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"-> {new_name}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_dvswitch(self, connector, config_list):
        data = config_list[0] if config_list else {}
        switches = data.get('switches', [])
        portgroups = data.get('portgroups', [])
        success_count = 0
        fail_count = 0

        manager = DVSwitchManager(connector)

        if switches:
            self._add_result(len(self.tree.get_children()) + 1, "开始", "创建分布式交换机...", "info")
            for sw in switches:
                success, msg = manager.create_dvswitch(
                    sw['datacenter'],
                    sw['switch_name'],
                    num_uplinks=4
                )
                self._add_result(len(self.tree.get_children()) + 1, "成功" if success else "失败",
                               f"交换机 {sw['switch_name']}: {msg}", "success" if success else "failure")
                if success:
                    success_count += 1
                else:
                    fail_count += 1

        if portgroups:
            self._add_result(len(self.tree.get_children()) + 1, "开始", "创建分布式端口组...", "info")
            for pg in portgroups:
                success, msg = manager.create_portgroup(
                    pg['switch_name'],
                    pg['portgroup_name'],
                    pg['vlan_id']
                )
                self._add_result(len(self.tree.get_children()) + 1, "成功" if success else "失败",
                               f"端口组 {pg['portgroup_name']}: {msg}", "success" if success else "failure")
                if success:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_cert(self, connector, config_list):
        manager = ESXiHostManager(connector)
        success_count = 0
        fail_count = 0

        def cert_single(config):
            success, msg = manager.update_certificate(config['esxi_host'])
            return (config['esxi_host'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(cert_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                host, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{host}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_ssh(self, connector, config_list):
        manager = ESXiHostManager(connector)
        success_count = 0
        fail_count = 0

        def ssh_single(config):
            success, msg = manager.enable_ssh(config['esxi_host'])
            return (config['esxi_host'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(ssh_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                host, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{host}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_ntp(self, connector, config_list):
        manager = ESXiHostManager(connector)
        success_count = 0
        fail_count = 0

        def ntp_single(config):
            success, msg = manager.config_ntp(config['esxi_host'], config['ntp_servers'])
            return (config['esxi_host'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(ntp_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                host, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{host}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_add_disk(self, connector, config_list):
        manager = VMDiskManager(connector)
        success_count = 0
        fail_count = 0

        def disk_single(config):
            vm_name = config['vm_name']
            disk_sizes = config['disk_sizes']
            datastore = config['datastore']
            results = manager.add_multiple_disks(vm_name, disk_sizes, datastore)
            ok = all(r[1] for r in results)
            msg = "; ".join([f"{r[0]}GB:{r[2]}" for r in results])
            return (vm_name, ok, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(disk_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                vm, ok, msg = future.result()
                self._add_result(len(self.tree.get_children()) + 1, "成功" if ok else "失败",
                               f"{vm}: {msg}", "success" if ok else "failure")
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _execute_resize(self, connector, config_list):
        manager = VMResizeManager(connector)
        success_count = 0
        fail_count = 0

        def resize_single(config):
            success, msg = manager.resize_vm_resources(
                vm_name=config['vm_name'],
                cpu_count=config.get('cpu_count'),
                memory_gb=config.get('memory_gb')
            )
            return (config['vm_name'], success, msg)

        with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
            futures = [executor.submit(resize_single, config) for config in config_list]
            for future in concurrent.futures.as_completed(futures):
                vm_name, ok, msg = future.result()
                self._add_result(
                    len(self.tree.get_children()) + 1,
                    "成功" if ok else "失败",
                    f"{vm_name}: {msg}",
                    "success" if ok else "failure"
                )
                if ok:
                    success_count += 1
                else:
                    fail_count += 1

        return success_count, fail_count

    def _parse_optional_positive_int(self, value):
        if value is None:
            return None
        text = str(value).strip()
        if text == '':
            return None
        try:
            number = int(float(text))
            return number if number > 0 else None
        except Exception:
            return None

    def _add_result(self, index, status, detail, tag='info'):
        self.tree.insert('', tk.END, values=(index, status, detail), tags=(tag,))
