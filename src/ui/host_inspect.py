import tkinter as tk
from tkinter import ttk, messagebox
import threading
import time
import pandas as pd
import concurrent.futures
from collections import defaultdict
from src.db.database import Database
from src.core.vcenter_connector import VCenterConnector
from src.core.host_collector import HostCollector
from src.ui.theme import COLORS, FONTS
from src.ui.enhanced_widgets import EnhancedButton, FilterableTable, ProgressDialog
from src.core.local_disk_inspector import LocalDiskInspector
from src.core.iscsi_manager import ISCSIManager
from src.core.host_nic_exporter import HostNicExporter
from src.core.storage_inspector import StorageInspector
from src.core.vm_collector import VMCollector
from src.core.excel_formatter import format_excel_file, highlight_column_cells


class HostInspectPage:
    DISPLAY_COLUMNS = ('vCenter', '主机名', '集群', 'CPU使用率', '内存使用率', '证书是否即将过期')
    COLUMN_WIDTHS = {
        'vCenter': 120, '主机名': 180, '集群': 120,
        'CPU使用率': 80, '内存使用率': 80, '证书是否即将过期': 150
    }
    DB_COLUMN_INDICES = [1, 2, 3, 15, 16, 17]
    DB_VC_IP_INDEX = 1
    DB_HOST_NAME_INDEX = 2

    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.db = Database()
        self.vc_list = {}
        self.all_data = []
        self._host_info_map = {}
        self._create_ui()

    def _create_ui(self):
        toolbar = ttk.Frame(self.parent, style='Card.TFrame')
        toolbar.pack(fill=tk.X, padx=10, pady=5)

        vc_frame = ttk.Frame(toolbar)
        vc_frame.pack(side=tk.LEFT)

        ttk.Label(vc_frame, text="vCenter:").pack(side=tk.LEFT)
        self.vc_combo = ttk.Combobox(vc_frame, width=20, state='readonly')
        self.vc_combo.pack(side=tk.LEFT, padx=5)
        self.vc_combo.bind('<<ComboboxSelected>>', lambda e: self._on_vc_selected())

        btn_frame = ttk.Frame(toolbar)
        btn_frame.pack(side=tk.LEFT, padx=10)

        EnhancedButton(btn_frame, "开始巡检", "🔍", command=self._start_inspect).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "一键巡检", "🚀", command=self._start_all_inspect).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "导出Excel", "📊", command=self._export_excel).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "iSCSI", "💾", command=self._export_iscsi).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "端口组", "🔀", command=self._export_portgroup).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "本地硬盘", "💿", command=self._start_local_disk_inspect_all).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "存储巡检", "🗄️", command=self._start_storage_inspect_all).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "存储巡检-跳过不可访问存储", "🗄️", command=self._start_storage_inspect_skip_inaccessible).pack(side=tk.LEFT, padx=2)
        EnhancedButton(btn_frame, "ESXi主机网卡巡检", "🌐", command=self._export_host_nics).pack(side=tk.LEFT, padx=2)

        log_frame = ttk.Frame(self.parent, style='Card.TFrame')
        log_frame.pack(fill=tk.X, padx=10, pady=(0, 5))
        ttk.Label(log_frame, text="日志:").pack(side=tk.LEFT, padx=5)
        self.log_text = tk.Text(log_frame, height=4, font=FONTS['small'])
        self.log_text.pack(fill=tk.X, padx=5, pady=5)

        self.filterable_tree = FilterableTable(
            self.parent,
            columns=self.DISPLAY_COLUMNS,
            column_widths=self.COLUMN_WIDTHS,
            page_size=100000,
            filter_columns=('vCenter', '集群', '证书是否即将过期')
        )
        self.filterable_tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        self.filterable_tree.bind_context_menu(self._show_context_menu)

        self._load_vc_list()
        self._load_all_data()

        self.parent.bind('<Visibility>', lambda e: self._refresh_vc_list())

    def _show_context_menu(self, event):
        self._log("右键菜单: 检测到点击事件")
        clicked_item = self.filterable_tree.tree.identify_row(event.y)
        if clicked_item:
            current_selection = self.filterable_tree.tree.selection()
            if clicked_item not in current_selection:
                self.filterable_tree.tree.selection_set(clicked_item)
        
        selected = self.filterable_tree.get_all_selected()
        self._log(f"右键菜单: 选中项目数={len(selected) if selected else 0}")
        if not selected:
            return

        count = len(selected)
        menu = tk.Menu(self.parent, tearoff=0)
        menu.add_command(label=f"iSCSI挂载 ({count}台)", command=lambda: self._iscsi_mount_batch(selected))
        menu.add_separator()
        menu.add_command(label=f"ESXi证书续订 ({count}台)", command=lambda: self._batch_cert(selected))
        menu.add_command(label=f"开启SSH ({count}台)", command=lambda: self._batch_ssh(selected))
        menu.add_command(label=f"配置NTP ({count}台)", command=lambda: self._batch_ntp(selected))
        menu.add_command(label=f"修改root密码 ({count}台)", command=lambda: self._batch_root_password(selected))
        menu.post(event.x_root, event.y_root)
        self._log(f"右键菜单: 已弹出菜单，处理{count}台主机")

    def _batch_cert(self, selected):
        self._log("_batch_cert called")
        self._execute_esxi_operation_direct(selected, 'cert')

    def _batch_ssh(self, selected):
        self._log("_batch_ssh called")
        self._execute_esxi_operation_direct(selected, 'ssh')

    def _batch_ntp(self, selected):
        self._log("_batch_ntp called")
        ntp_server = self._ask_ntp_server()
        if ntp_server:
            self._execute_esxi_operation_direct(selected, 'ntp', ntp_server)

    def _batch_root_password(self, selected):
        self._log("_batch_root_password called")
        new_password = self._ask_root_password()
        if new_password:
            self._execute_esxi_operation_direct(selected, 'root_password', root_password=new_password)

    def _ask_ntp_server(self):
        dialog = tk.Toplevel(self.parent)
        dialog.title("输入NTP服务器")
        dialog.geometry("400x150")
        dialog.transient(self.parent)
        dialog.grab_set()

        result = [None]

        ttk.Label(dialog, text="NTP服务器地址(空格分隔多个):").pack(pady=10)
        ntp_var = tk.StringVar()
        ntp_entry = ttk.Entry(dialog, textvariable=ntp_var, width=40)
        ntp_entry.pack(pady=5)

        def on_ok():
            result[0] = ntp_var.get()
            dialog.destroy()

        def on_cancel():
            dialog.destroy()

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="确定", command=on_ok).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="取消", command=on_cancel).pack(side=tk.LEFT, padx=5)

        dialog.wait_window()
        return result[0]

    def _ask_root_password(self):
        dialog = tk.Toplevel(self.parent)
        dialog.title("输入root新密码")
        dialog.geometry("420x230")
        dialog.transient(self.parent)
        dialog.grab_set()

        result = [None]

        ttk.Label(dialog, text="root新密码:").pack(pady=(15, 6))
        pwd_var = tk.StringVar()
        pwd_entry = ttk.Entry(dialog, textvariable=pwd_var, width=40, show='*')
        pwd_entry.pack(pady=4)
        pwd_entry.focus()

        ttk.Label(dialog, text="确认root新密码:").pack(pady=(8, 6))
        confirm_var = tk.StringVar()
        confirm_entry = ttk.Entry(dialog, textvariable=confirm_var, width=40, show='*')
        confirm_entry.pack(pady=4)

        ttk.Label(dialog, text="两次密码必须一致；该操作会按选中主机并行执行", foreground=COLORS['text_secondary']).pack(pady=(6, 0))

        def on_ok():
            value = pwd_var.get().strip()
            confirm_value = confirm_var.get().strip()
            if not value:
                messagebox.showwarning("提示", "请输入root新密码")
                return
            if value != confirm_value:
                messagebox.showwarning("提示", "两次输入的密码不一致")
                return
            result[0] = value
            dialog.destroy()

        def on_cancel():
            dialog.destroy()

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=12)
        ttk.Button(btn_frame, text="确定", command=on_ok).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="取消", command=on_cancel).pack(side=tk.LEFT, padx=5)

        dialog.wait_window()
        return result[0]

    def _execute_esxi_operation_direct(self, selected, operation, ntp_server=None, root_password=None):
        host_info_list = []
        for item in selected:
            item_values = item
            if len(item_values) >= 2:
                vc_ip = item_values[0]
                host_name = item_values[1]
                host_info_list.append((vc_ip, host_name))

        if not host_info_list:
            messagebox.showwarning("提示", "无法获取选中主机信息")
            return

        self._log(f"开始执行{operation}操作... 共{len(host_info_list)}台")
        progress = ProgressDialog(self.parent, f"执行{operation}操作")

        def run_operation():
            from src.core.esxi_host_manager import ESXiHostManager
            from src.core.esxi_root_password_manager import ESXiRootPasswordManager
            success_count = 0
            fail_count = 0
            total = len(host_info_list)
            lock = threading.Lock()
            processed = [0]

            vc_groups = defaultdict(list)
            for vc_ip, host_name in host_info_list:
                vc_groups[vc_ip].append(host_name)

            def process_single_host(vc_ip, host_name):
                user, pwd = self.vc_list.get(vc_ip, (None, None))
                if not user or not pwd:
                    with lock:
                        processed[0] += 1
                        fail_count_inc = fail_count + 1
                    self._log(f"VC {vc_ip} 连接信息不存在")
                    return (host_name, False, "VC连接信息不存在", None)

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()
                if not success:
                    connector.disconnect()
                    return (host_name, False, f"连接失败: {msg}", None)

                try:
                    manager = ESXiHostManager(connector)
                    pwd_manager = ESXiRootPasswordManager(connector)
                    if operation == 'cert':
                        ok, msg = manager.update_certificate(host_name)
                    elif operation == 'ssh':
                        ok, msg = manager.enable_ssh(host_name)
                    elif operation == 'ntp':
                        ok, msg = manager.config_ntp(host_name, ntp_server)
                    elif operation == 'root_password':
                        ok, msg = pwd_manager.change_root_password(host_name, root_password)
                    else:
                        ok, msg = False, "未知操作"
                    return (host_name, ok, msg, None)
                except Exception as e:
                    return (host_name, False, None, str(e))
                finally:
                    connector.disconnect()

            max_workers = min(15, len(host_info_list)) or 1
            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = []
                for vc_ip, host_names in vc_groups.items():
                    for host_name in host_names:
                        future = executor.submit(process_single_host, vc_ip, host_name)
                        futures.append(future)
                
                for future in concurrent.futures.as_completed(futures):
                    try:
                        host_name, ok, msg, error = future.result()
                        with lock:
                            processed[0] += 1
                        if error:
                            self._log(f"[{processed[0]}/{total}] {host_name}: 异常 {error}")
                            with lock:
                                fail_count += 1
                        else:
                            self._log(f"[{processed[0]}/{total}] {host_name}: {msg}")
                            with lock:
                                if ok:
                                    success_count += 1
                                else:
                                    fail_count += 1
                    except Exception as e:
                        self._log(f"处理异常: {str(e)}")

            progress.close()
            self._log(f"{operation}完成: 成功={success_count}, 失败={fail_count}")
            self.parent.after(0, lambda: messagebox.showinfo("完成", f"成功: {success_count}, 失败: {fail_count}"))

        threading.Thread(target=run_operation, daemon=True).start()

    def _load_vc_list(self):
        vcenters = self.db.get_vcenters()
        self.vc_list = {vc[1]: (vc[2], vc[3]) for vc in vcenters}
        vc_options = ['全部'] + list(self.vc_list.keys())
        self.vc_combo['values'] = vc_options
        if vc_options:
            self.vc_combo.current(0)

    def _refresh_vc_list(self):
        vcenters = self.db.get_vcenters()
        self.vc_list = {vc[1]: (vc[2], vc[3]) for vc in vcenters}
        current = self.vc_combo.get()
        vc_options = ['全部'] + list(self.vc_list.keys())
        self.vc_combo['values'] = vc_options
        if current in vc_options:
            self.vc_combo.set(current)
        else:
            self.vc_combo.current(0)
            self._load_all_data()

    def _on_vc_selected(self):
        self._load_all_data()

    def _log(self, message):
        self.log_text.insert(tk.END, f"{message}\n")
        self.log_text.see(tk.END)

    def _load_all_data(self):
        selected_vc = self.vc_combo.get()
        if selected_vc == '全部':
            data, _ = self.db.get_host_inspect(None, '', 1, 100000)
        else:
            data, _ = self.db.get_host_inspect(selected_vc, '', 1, 100000)

        self._host_info_map = {}
        filtered_data = []
        for row in data:
            new_row = tuple(row[i] for i in self.DB_COLUMN_INDICES)
            new_row = tuple('' if v is None else v for v in new_row)
            filtered_data.append(new_row)
            vc_ip = row[self.DB_VC_IP_INDEX] if self.DB_VC_IP_INDEX < len(row) else None
            host_name = row[self.DB_HOST_NAME_INDEX] if self.DB_HOST_NAME_INDEX < len(row) else None
            if vc_ip and host_name:
                self._host_info_map[new_row] = (vc_ip, host_name)

        self.all_data = filtered_data
        self.filterable_tree.set_data(filtered_data)

    def _iscsi_mount_batch(self, selected_items):
        if not selected_items:
            messagebox.showwarning("提示", "请先选择主机")
            return

        host_info_list = []
        for item in selected_items:
            item_values = item
            if len(item_values) >= 2:
                vc_ip = str(item_values[0]).strip()
                host_name = str(item_values[1]).strip()
                if vc_ip and host_name:
                    host_info_list.append((vc_ip, host_name))

        if not host_info_list:
            messagebox.showwarning("提示", "无法获取选中主机信息")
            return

        # 去重，避免重复选择导致同一主机重复执行
        host_info_list = list(dict.fromkeys(host_info_list))

        dialog = tk.Toplevel(self.parent)
        dialog.title(f"批量iSCSI挂载 ({len(host_info_list)}台)")
        dialog.geometry("500x180")
        dialog.transient(self.parent)
        dialog.grab_set()

        tk.Label(dialog, text=f"已选择 {len(host_info_list)} 台主机", font=FONTS['bold']).pack(pady=(10, 5))

        form_frame = tk.Frame(dialog)
        form_frame.pack(fill=tk.X, padx=20, pady=5)

        tk.Label(form_frame, text="iSCSI动态发现IP:").grid(row=0, column=0, sticky='w', pady=5)
        target_entry = tk.Entry(form_frame, width=30, font=FONTS['normal'])
        target_entry.grid(row=0, column=1, padx=5, pady=5)

        def do_mount():
            target_ip = target_entry.get().strip()
            if not target_ip:
                messagebox.showwarning("提示", "请输入iSCSI Target IP")
                return
            target_ips = [ip.strip() for ip in target_ip.replace('\n', ',').replace(' ', ',').split(',') if ip.strip()]
            if not target_ips:
                messagebox.showwarning("提示", "请输入有效的iSCSI Target IP")
                return
            dialog.destroy()
            progress = ProgressDialog(self.parent, "批量iSCSI挂载")
            progress.update_status("开始批量挂载...")

            def run_mount():
                try:
                    success_count = 0
                    fail_count = 0
                    total = len(host_info_list)
                    lock = threading.Lock()
                    processed = [0]

                    def mount_single(vc_ip, host_name):
                        vc_user, vc_pwd = self.vc_list.get(vc_ip, (None, None))
                        if not vc_user or not vc_pwd:
                            return (host_name, False, "VC连接信息不存在")

                        connector = VCenterConnector(vc_ip, vc_user, vc_pwd)
                        connected, connect_msg = connector.connect()
                        if not connected:
                            connector.disconnect()
                            return (host_name, False, f"连接失败: {connect_msg}")

                        manager = ISCSIManager(connector)
                        try:
                            ok, msg = manager.add_iscsi_target(
                                esxi_host_name=host_name,
                                target_ips=target_ips
                            )
                            return (host_name, ok, msg)
                        except Exception as e:
                            return (host_name, False, str(e))
                        finally:
                            connector.disconnect()

                    max_workers = min(15, len(host_info_list)) or 1
                    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                        futures = [executor.submit(mount_single, vc_ip, host_name) 
                                   for vc_ip, host_name in host_info_list]
                        
                        for future in concurrent.futures.as_completed(futures):
                            try:
                                host_name, ok, msg = future.result()
                                with lock:
                                    processed[0] += 1
                                progress.update_status(f"正在处理 {processed[0]}/{total}: {host_name}")
                                self._log(f"[{processed[0]}/{total}] {host_name}: {msg}")
                                with lock:
                                    if ok:
                                        success_count += 1
                                    else:
                                        fail_count += 1
                            except Exception as e:
                                self._log(f"处理异常: {str(e)}")

                    progress.close()
                    result_msg = f"批量挂载完成\n成功: {success_count}台\n失败: {fail_count}台"
                    self.parent.after(0, lambda: messagebox.showinfo("完成", result_msg))
                except Exception as e:
                    progress.close()
                    self.parent.after(0, lambda: messagebox.showerror("错误", f"挂载失败: {str(e)}"))

            threading.Thread(target=run_mount, daemon=True).start()

        btn_frame = tk.Frame(dialog)
        btn_frame.pack(pady=10)
        tk.Button(btn_frame, text="挂载", command=do_mount, font=FONTS['button'], bg=COLORS['primary'], fg='white').pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="取消", command=dialog.destroy, font=FONTS['button']).pack(side=tk.LEFT, padx=5)

    def _start_inspect(self):
        vc_ip = self.vc_combo.get()
        if not vc_ip or vc_ip == '全部':
            messagebox.showwarning("提示", "请选择具体的vCenter")
            return

        user, pwd = self.vc_list[vc_ip]
        progress = ProgressDialog(self.parent, "ESXi主机巡检中")

        def run_inspect():
            try:
                start_time = time.time()
                progress.update_status(f"正在连接 {vc_ip}...")
                self._log(f"开始巡检 {vc_ip}...")

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()

                if not success:
                    progress.close()
                    self.parent.after(0, lambda: messagebox.showerror("错误", f"连接失败: {msg}"))
                    self._log(f"连接 {vc_ip} 失败: {msg}")
                    return

                progress.update_status("正在采集主机数据...")
                self._log(f"正在采集 {vc_ip} 主机数据...")

                collector = HostCollector(connector)
                hosts = collector.collect_all_hosts()

                self.db.save_host_inspect(hosts)
                connector.disconnect()

                elapsed = time.time() - start_time
                progress.close()

                self._log(f"巡检完成: {vc_ip} 共 {len(hosts)} 台主机，耗时 {elapsed:.1f}秒")
                self.parent.after(0, lambda: messagebox.showinfo("完成", f"巡检完成，共采集 {len(hosts)} 台主机\n耗时: {elapsed:.1f}秒"))
                self.parent.after(0, lambda: self._load_all_data())

            except Exception as e:
                progress.close()
                self._log(f"巡检失败: {str(e)}")
                self.parent.after(0, lambda: messagebox.showerror("错误", f"巡检失败: {str(e)}"))

        threading.Thread(target=run_inspect, daemon=True).start()

    def _start_storage_inspect_all(self):
        self._start_storage_inspect_common(skip_inaccessible=False, title_suffix='')

    def _start_storage_inspect_skip_inaccessible(self):
        self._start_storage_inspect_common(skip_inaccessible=True, title_suffix='(跳过不可访问)')

    def _start_storage_inspect_common(self, skip_inaccessible=False, title_suffix=''):
        from tkinter import filedialog
        import datetime

        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存存储巡检结果",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile=f"存储巡检{title_suffix}_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
        )
        if not file_path:
            return

        progress = ProgressDialog(self.parent, "存储巡检中")
        region_map = self.db.get_vcenter_region_map()

        def inspect_single_vc(vc):
            vc_ip, user, pwd = vc[1], vc[2], vc[3]
            connector = VCenterConnector(vc_ip, user, pwd)
            success, msg = connector.connect()
            if not success:
                raise RuntimeError(f"{vc_ip}: {msg}")
            try:
                inspector = StorageInspector(connector)
                storage_rows = inspector.collect_all_datastores(skip_inaccessible=skip_inaccessible)
                return storage_rows
            finally:
                connector.disconnect()

        def run_inspect():
            try:
                from src.core.concurrent_config import ConcurrentConfig
                all_rows = []
                failed = 0
                progress.update_status("正在收集存储信息...")
                workers = min(ConcurrentConfig.get_optimal_workers('vc_inspect'), len(vcenters))
                with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
                    futures = {executor.submit(inspect_single_vc, vc): vc for vc in vcenters}
                    for idx, future in enumerate(concurrent.futures.as_completed(futures), start=1):
                        progress.update_status(f"正在执行存储巡检... {idx}/{len(vcenters)}")
                        try:
                            storage_rows = future.result()
                            all_rows.extend(storage_rows)
                        except Exception as e:
                            failed += 1
                            self._log(f"存储巡检失败: {e}")

                progress.close()
                if not all_rows:
                    self.parent.after(0, lambda: messagebox.showwarning("提示", "未采集到存储信息"))
                    return

                StorageInspector.save_to_excel(file_path, all_rows, region_map=region_map)
                format_excel_file(file_path)
                msg = f"存储巡检完成，共 {len(all_rows)} 条记录\n文件: {file_path}"
                if failed:
                    msg += f"\n失败: {failed} 个vCenter"
                self.parent.after(0, lambda: messagebox.showinfo("完成", msg))
            except Exception as e:
                progress.close()
                self.parent.after(0, lambda: messagebox.showerror("错误", f"巡检失败: {e}"))

        threading.Thread(target=run_inspect, daemon=True).start()

    def _start_all_inspect(self):
        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        self._log("开始一键巡检...")
        progress = ProgressDialog(self.parent, "一键巡检中")

        start_time = time.time()
        completed = [0]
        total_hosts = [0]
        results = []
        lock = threading.Lock()

        def inspect_single_vc(vc):
            vc_ip, user, pwd = vc[1], vc[2], vc[3]
            try:
                self.parent.after(0, lambda ip=vc_ip: self._log(f"正在连接 {ip}..."))
                progress.update_status(f"正在连接 {vc_ip}...")

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()

                if not success:
                    self.parent.after(0, lambda ip=vc_ip, m=msg: self._log(f"连接 {ip} 失败: {m}"))
                    return {'vc_ip': vc_ip, 'success': False, 'error': msg, 'hosts': 0}

                self.parent.after(0, lambda ip=vc_ip: self._log(f"正在采集 {ip} 主机数据..."))
                progress.update_status(f"正在采集 {vc_ip} 主机数据...")
                collector = HostCollector(connector)
                hosts = collector.collect_all_hosts()
                self.db.save_host_inspect(hosts)
                connector.disconnect()

                self.parent.after(0, lambda ip=vc_ip, h=len(hosts): self._log(f"{ip} 采集完成，共 {h} 台主机"))

                with lock:
                    completed[0] += 1
                    progress.update_detail(f"已完成 {completed[0]}/{len(vcenters)} 个vCenter")

                return {'vc_ip': vc_ip, 'success': True, 'hosts': len(hosts), 'error': None}

            except Exception as e:
                self.parent.after(0, lambda ip=vc_ip, err=str(e): self._log(f"{ip} 失败: {err}"))
                with lock:
                    completed[0] += 1
                return {'vc_ip': vc_ip, 'success': False, 'error': str(e), 'hosts': 0}

        def run_inspect():
            from concurrent.futures import ThreadPoolExecutor, as_completed
            from src.core.concurrent_config import ConcurrentConfig

            max_workers = ConcurrentConfig.get_optimal_workers('vc_inspect')
            max_workers = min(max_workers, len(vcenters))

            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {executor.submit(inspect_single_vc, vc): vc for vc in vcenters}

                for future in as_completed(futures):
                    result = future.result()
                    results.append(result)
                    total_hosts[0] += result['hosts']

            elapsed = int(time.time() - start_time)
            minutes = elapsed // 60
            seconds = elapsed % 60
            time_str = f"{minutes}分{seconds}秒" if minutes > 0 else f"{seconds}秒"

            failed = sum(1 for r in results if not r['success'])

            progress.close()

            msg = f"一键巡检完成\n总采集: {total_hosts[0]} 台主机\n耗时: {time_str}"
            if failed > 0:
                msg += f"\n失败: {failed} 个vCenter"

            self._log(msg)
            self.parent.after(0, lambda: messagebox.showinfo("完成", msg))
            self.parent.after(0, lambda: self._on_vc_selected())

        threading.Thread(target=run_inspect, daemon=True).start()

    def _export_excel(self):
        from tkinter import filedialog

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")]
        )

        if not file_path:
            return

        try:
            selected_vc = self.vc_combo.get()
            self._log(f"正在导出 {'全部' if selected_vc == '全部' else selected_vc} 主机数据...")
            region_map = self.db.get_vcenter_region_map()

            if selected_vc == '全部':
                data, _ = self.db.get_host_inspect(None, '', 1, 100000)
            else:
                data, _ = self.db.get_host_inspect(selected_vc, '', 1, 100000)

            if data:
                df = pd.DataFrame(data, columns=[
                    'ID', 'vCenter', '主机名', '集群', '状态', 'ESXi版本', '厂商', '型号',
                    '序列号', 'CPU型号', '主频', 'CPU数', '核心数', '内存(GB)', '运行天数',
                    'CPU使用率', '内存使用率', '证书是否即将过期', 'VMK适配器', '证书过期', '巡检时间'
                ])
                df.insert(0, '区域', df['vCenter'].map(lambda x: region_map.get(str(x), '') if pd.notna(x) else ''))
                df.to_excel(file_path, index=False)
                format_excel_file(file_path)
                highlight_column_cells(file_path, '证书是否即将过期', ['是'])
                self._log(f"导出完成: {len(data)} 条记录")
                messagebox.showinfo("提示", f"导出成功: {file_path}")
            else:
                self._log("没有数据可导出")
                messagebox.showwarning("提示", "没有数据")
        except Exception as e:
            self._log(f"导出失败: {str(e)}")
            messagebox.showerror("错误", f"导出失败: {str(e)}")

    def _export_iscsi(self):
        from tkinter import filedialog
        from src.core.iscsi_exporter import ISCSIExporter

        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存iSCSI信息",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile="esxi_iscsi_iqn.xlsx"
        )

        if not file_path:
            return

        progress = ProgressDialog(self.parent, "导出iSCSI信息")

        def run_export():
            try:
                start_time = time.time()
                self._log("开始导出iSCSI信息...")
                exporter = ISCSIExporter()
                results = exporter.export_iscsi(
                    vcenters,
                    status_callback=lambda msg: progress.update_status(msg)
                )

                elapsed = time.time() - start_time

                if results:
                    exporter.save_to_excel(file_path)
                    progress.close()
                    self._log(f"导出完成: {len(results)} 条记录，耗时 {elapsed:.1f}秒")
                    self.parent.after(0, lambda: messagebox.showinfo("完成", f"导出成功，共 {len(results)} 条记录\n耗时: {elapsed:.1f}秒"))
                else:
                    progress.close()
                    self._log("未获取到iSCSI信息")
                    self.parent.after(0, lambda: messagebox.showwarning("提示", "未获取到iSCSI信息"))

            except Exception as e:
                progress.close()
                self._log(f"导出失败: {str(e)}")
                self.parent.after(0, lambda: messagebox.showerror("错误", f"导出失败: {str(e)}"))

        threading.Thread(target=run_export, daemon=True).start()

    def _export_portgroup(self):
        from tkinter import filedialog
        import traceback
        from src.core.portgroup_exporter import PortGroupExporter

        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存端口组信息",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile="portgroup_info.xlsx"
        )

        if not file_path:
            return

        progress = ProgressDialog(self.parent, "导出端口组信息")

        def run_export():
            try:
                self._log("开始导出端口组信息...")
                exporter = PortGroupExporter()

                def update_status(msg):
                    progress.update_status(msg)

                progress.update_status(f"开始收集 {len(vcenters)} 个vCenter...")

                results, total_time = exporter.export_all_vcenters(vcenters, status_callback=update_status)

                progress.update_status(f"收集完成，共 {len(results)} 个端口组 (耗时: {total_time:.1f}s)")

                if results:
                    progress.update_status("正在写入Excel...")
                    exporter.save_to_excel(file_path, results)
                    progress.close()
                    self._log(f"导出完成: {len(results)} 条记录，耗时 {total_time:.1f}秒")
                    self.parent.after(0, lambda: messagebox.showinfo("完成", f"导出成功，共 {len(results)} 条记录"))
                else:
                    progress.close()
                    self._log("未获取到端口组信息")
                    self.parent.after(0, lambda: messagebox.showwarning("提示", "未获取到端口组信息"))

            except Exception as e:
                error_msg = f"导出失败: {str(e)}"
                print(traceback.format_exc())
                progress.close()
                self._log(error_msg)
                self.parent.after(0, lambda: messagebox.showerror("错误", error_msg))

        threading.Thread(target=run_export, daemon=True).start()

    def _export_host_nics(self):
        from tkinter import filedialog

        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存ESXi主机网卡巡检结果",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile=f"ESXi主机网卡巡检_{time.strftime('%m%d')}.xlsx"
        )

        if not file_path:
            return

        progress = ProgressDialog(self.parent, "ESXi主机网卡巡检中")

        def run_export():
            try:
                self._log("开始执行 ESXi 主机网卡巡检...")
                exporter = HostNicExporter()
                rows = exporter.export_all_vcenters(
                    vcenters,
                    status_callback=lambda msg: progress.update_status(msg)
                )

                if not rows:
                    progress.close()
                    self._log("未获取到ESXi主机网卡信息")
                    self.parent.after(0, lambda: messagebox.showwarning("提示", "未获取到ESXi主机网卡信息"))
                    return

                exporter.save_to_excel(file_path, rows)
                progress.close()
                self._log(f"ESXi主机网卡巡检完成，共 {len(rows)} 条记录")
                self.parent.after(0, lambda: messagebox.showinfo("完成", f"ESXi主机网卡巡检完成\n共 {len(rows)} 条记录\n文件: {file_path}"))
            except Exception as e:
                progress.close()
                self._log(f"ESXi主机网卡巡检失败: {str(e)}")
                self.parent.after(0, lambda: messagebox.showerror("错误", f"ESXi主机网卡巡检失败: {str(e)}"))

        threading.Thread(target=run_export, daemon=True).start()

    def _start_local_disk_inspect_one(self):
        vc_ip = self.vc_combo.get()
        if not vc_ip or vc_ip == '全部':
            messagebox.showwarning("提示", "请选择具体的vCenter")
            return

        user, pwd = self.vc_list[vc_ip]
        progress = ProgressDialog(self.parent, "本地硬盘巡检中")

        def run_inspect():
            try:
                start_time = time.time()
                progress.update_status(f"正在连接 {vc_ip}...")

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()

                if not success:
                    progress.close()
                    self.parent.after(0, lambda: messagebox.showerror("错误", f"连接失败: {msg}"))
                    return

                progress.update_status("正在采集本地硬盘数据...")

                inspector = LocalDiskInspector(connector)
                disks = inspector.collect_all_local_disks(
                    progress_callback=lambda x: progress.update_status(f"已处理 {x} 台主机...")
                )

                progress.update_status("正在保存数据...")
                self.db.save_local_disk_inspect(disks)

                connector.disconnect()

                elapsed = time.time() - start_time
                progress.close()

                self.parent.after(0, lambda: messagebox.showinfo("完成", f"本地硬盘巡检完成，共采集 {len(disks)} 块硬盘\n耗时: {elapsed:.1f}秒"))

            except Exception as e:
                progress.close()
                self.parent.after(0, lambda: messagebox.showerror("错误", f"巡检失败: {str(e)}"))

        threading.Thread(target=run_inspect, daemon=True).start()

    def _start_local_disk_inspect_all(self):
        from tkinter import filedialog
        import datetime
        import os

        vcenters = self.db.get_vcenters()
        if not vcenters:
            messagebox.showwarning("提示", "没有已添加的vCenter")
            return

        file_path = filedialog.asksaveasfilename(
            parent=self.parent,
            title="保存本地硬盘巡检",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile=f"ESXi本地硬盘巡检_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
        )

        if not file_path:
            return

        self._log("开始本地硬盘一键巡检...")
        progress = ProgressDialog(self.parent, "本地硬盘一键巡检中")

        start_time = time.time()
        completed = [0]
        total_disks = [0]
        results = []
        all_disk_info = []
        lock = threading.Lock()

        def inspect_single_vc(vc):
            vc_ip, user, pwd = vc[1], vc[2], vc[3]
            try:
                self.parent.after(0, lambda ip=vc_ip: self._log(f"正在连接 {ip}..."))
                progress.update_status(f"正在连接 {vc_ip}...")

                connector = VCenterConnector(vc_ip, user, pwd)
                success, msg = connector.connect()

                if not success:
                    self.parent.after(0, lambda ip=vc_ip, m=msg: self._log(f"连接 {ip} 失败: {m}"))
                    return {'vc_ip': vc_ip, 'success': False, 'error': msg, 'disks': 0, 'disk_info': []}

                self.parent.after(0, lambda ip=vc_ip: self._log(f"正在采集 {ip} 本地硬盘数据..."))
                progress.update_status(f"正在采集 {vc_ip} 本地硬盘数据...")
                inspector = LocalDiskInspector(connector)
                disks = inspector.collect_all_local_disks()
                self.db.save_local_disk_inspect(disks)
                connector.disconnect()

                self.parent.after(0, lambda ip=vc_ip, d=len(disks): self._log(f"{ip} 采集完成，共 {d} 块硬盘"))

                with lock:
                    completed[0] += 1
                    progress.update_detail(f"已完成 {completed[0]}/{len(vcenters)} 个vCenter")

                return {'vc_ip': vc_ip, 'success': True, 'disks': len(disks), 'error': None, 'disk_info': disks}

            except Exception as e:
                self.parent.after(0, lambda ip=vc_ip, err=str(e): self._log(f"{ip} 失败: {err}"))
                with lock:
                    completed[0] += 1
                return {'vc_ip': vc_ip, 'success': False, 'error': str(e), 'disks': 0, 'disk_info': []}

        def run_inspect():
            from concurrent.futures import ThreadPoolExecutor, as_completed
            from src.core.concurrent_config import ConcurrentConfig
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, PatternFill, Font
            from collections import defaultdict

            max_workers = ConcurrentConfig.get_optimal_workers('vc_inspect')
            max_workers = min(max_workers, len(vcenters))

            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {executor.submit(inspect_single_vc, vc): vc for vc in vcenters}

                for future in as_completed(futures):
                    result = future.result()
                    results.append(result)
                    total_disks[0] += result['disks']
                    all_disk_info.extend(result['disk_info'])

            elapsed = int(time.time() - start_time)
            minutes = elapsed // 60
            seconds = elapsed % 60
            time_str = f"{minutes}分{seconds}秒" if minutes > 0 else f"{seconds}秒"

            failed = sum(1 for r in results if not r['success'])

            if not all_disk_info:
                progress.close()
                self._log("未采集到任何本地硬盘信息")
                self.parent.after(0, lambda: messagebox.showwarning("提示", "未采集到任何本地硬盘信息"))
                return

            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "ESXi本地硬盘巡检"
                region_map = self.db.get_vcenter_region_map()

                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

                headers = ['VC', '区域', '数据中心', '集群', 'ESXi主机名称', '硬盘名称', '所属数据存储', '硬盘总容量(GB)', '硬盘已用容量(GB)', '硬盘空间使用率']
                for col_num, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_num, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align

                host_groups = defaultdict(list)
                for disk_info in all_disk_info:
                    host_key = (disk_info['vc_ip'], disk_info['datacenter'], disk_info['cluster'], disk_info['host_name'])
                    host_groups[host_key].append(disk_info)

                current_row = 2
                merge_ranges = []

                for host_key, disks in host_groups.items():
                    vc_ip, datacenter, cluster, host_name = host_key
                    region = region_map.get(str(vc_ip), '')
                    num_disks = len(disks)

                    for i, disk in enumerate(disks):
                        row = current_row + i

                        ws.cell(row=row, column=1, value=disk['vc_ip'])
                        ws.cell(row=row, column=2, value=region)
                        ws.cell(row=row, column=3, value=disk['datacenter'])
                        ws.cell(row=row, column=4, value=disk['cluster'])
                        ws.cell(row=row, column=5, value=disk['host_name'])
                        ws.cell(row=row, column=6, value=disk['disk_name'])
                        ws.cell(row=row, column=7, value=disk['datastore'])
                        ws.cell(row=row, column=8, value=disk['total_gb'])
                        ws.cell(row=row, column=9, value=disk['used_gb'])
                        ws.cell(row=row, column=10, value=disk['usage_percent'])

                        for col in range(1, 11):
                            ws.cell(row=row, column=col).alignment = center_align

                    if num_disks > 1:
                        start_row = current_row
                        end_row = current_row + num_disks - 1
                        for col in range(1, 6):
                            merge_ranges.append((start_row, col, end_row, col))

                    current_row += num_disks

                for start_row, start_col, end_row, end_col in merge_ranges:
                    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
                    merged_cell = ws.cell(row=start_row, column=start_col)
                    merged_cell.alignment = center_align

                column_widths = {'A': 15, 'B': 12, 'C': 20, 'D': 20, 'E': 25, 'F': 50, 'G': 25, 'H': 18, 'I': 18, 'J': 20}
                for col_letter, width in column_widths.items():
                    ws.column_dimensions[col_letter].width = width

                ws.freeze_panes = 'A2'
                wb.save(file_path)
                format_excel_file(file_path)

                progress.close()

                msg = f"本地硬盘一键巡检完成\n总采集: {total_disks[0]} 块硬盘\n耗时: {time_str}\n文件: {file_path}"
                if failed > 0:
                    msg += f"\n失败: {failed} 个vCenter"

                self._log(msg)
                self.parent.after(0, lambda: messagebox.showinfo("完成", msg))

            except Exception as e:
                progress.close()
                self.parent.after(0, lambda err=str(e): messagebox.showerror("错误", f"导出失败: {err}"))

        threading.Thread(target=run_inspect, daemon=True).start()
