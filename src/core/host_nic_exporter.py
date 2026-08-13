import concurrent.futures
import re
import time
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.core.concurrent_config import ConcurrentConfig
from src.core.excel_formatter import format_excel_file
from src.core.vcenter_connector import VCenterConnector
from src.db.database import Database


class HostNicExporter:
    def __init__(self):
        self.db = Database()

    def export_all_vcenters(self, vcenters, status_callback=None):
        all_rows = []

        def collect_single(vc_config):
            vc_ip = vc_config[1]
            username = vc_config[2]
            password = vc_config[3]
            connector = VCenterConnector(vc_ip, username, password)
            success, msg = connector.connect()
            if not success:
                raise RuntimeError(f"{vc_ip}: {msg}")

            try:
                if status_callback:
                    status_callback(f"正在采集 {vc_ip} 网卡信息...")
                rows = self.collect_host_nics(connector)
                return rows, vc_ip
            finally:
                connector.disconnect()

        max_workers = min(ConcurrentConfig.get_optimal_workers("vc_inspect", len(vcenters)), len(vcenters)) or 1
        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {executor.submit(collect_single, vc): vc for vc in vcenters}
            for idx, future in enumerate(concurrent.futures.as_completed(futures), start=1):
                vc = futures[future]
                vc_ip = vc[1]
                if status_callback:
                    status_callback(f"正在执行 ESXi 主机网卡巡检... {idx}/{len(vcenters)}")
                try:
                    rows, _ = future.result()
                    all_rows.extend(rows)
                except Exception as exc:
                    if status_callback:
                        status_callback(f"{vc_ip} 采集失败: {exc}")

        all_rows.sort(key=lambda row: (
            str(row.get("vc_ip", "")),
            str(row.get("host_name", "")),
            self._nic_sort_key(str(row.get("device", ""))),
        ))
        return all_rows

    def collect_host_nics(self, connector):
        if not connector.is_connected():
            return []

        cluster_map = connector.build_cluster_host_map()
        result = connector.get_all_hosts()

        rows = []
        for obj_content in result:
            props = {p.name: p.val for p in obj_content.propSet}
            host_mor = obj_content.obj
            host_moid = getattr(host_mor, "_moId", "")
            host_name = props.get("name", "")
            cluster_name = cluster_map.get(host_moid, "")
            network_config = props.get("config.network")
            rows.extend(self._parse_host_nics(connector.ip, host_name, cluster_name, network_config))

        return rows

    def _parse_host_nics(self, vc_ip, host_name, cluster_name, network_config):
        if not network_config:
            return []

        switch_map = self._build_switch_map(network_config)
        pnics = list(getattr(network_config, "pnic", []) or [])
        rows = []

        for pnic in pnics:
            device = str(getattr(pnic, "device", "") or "").strip()
            if not device:
                continue

            actual_speed = self._get_actual_speed(pnic)
            configured_speed = self._get_configured_speed(pnic)
            switch_name = self._get_switch_name(device, pnic, switch_map)

            rows.append(
                {
                    "vc_ip": vc_ip,
                    "host_name": host_name,
                    "cluster": cluster_name,
                    "device": device,
                    "actual_speed": actual_speed,
                    "configured_speed": configured_speed,
                    "switch": switch_name,
                }
            )

        return rows

    def _build_switch_map(self, network_config):
        switch_map = defaultdict(set)

        try:
            for vswitch in list(getattr(network_config, "vswitch", []) or []):
                switch_name = str(getattr(vswitch, "name", "") or getattr(vswitch, "key", "") or "").strip()
                if not switch_name:
                    continue
                for pnic in list(getattr(vswitch, "pnic", []) or []):
                    device = str(pnic or "").strip()
                    if device:
                        switch_map[device].add(switch_name)
        except Exception:
            pass

        for proxy_switch in list(getattr(network_config, "proxySwitch", []) or []):
            switch_name = (
                str(getattr(proxy_switch, "name", "") or getattr(proxy_switch, "dvsName", "") or getattr(proxy_switch, "key", "") or "").strip()
            )
            if not switch_name:
                continue

            for device in self._extract_proxy_pnics(proxy_switch):
                if device:
                    switch_map[device].add(switch_name)

        return switch_map

    def _extract_proxy_pnics(self, proxy_switch):
        devices = set()
        for attr_name in ("pnicSpec", "pnic", "pnicKey"):
            value = getattr(proxy_switch, attr_name, None)
            devices.update(self._extract_devices(value))

        spec = getattr(proxy_switch, "spec", None)
        if spec is not None:
            for attr_name in ("backing", "pnicSpec", "pnic", "pnicKey"):
                value = getattr(spec, attr_name, None)
                devices.update(self._extract_devices(value))

        return devices

    def _extract_devices(self, value):
        devices = set()
        if not value:
            return devices

        if isinstance(value, str):
            if value.startswith("vmnic"):
                devices.add(value)
            return devices

        if isinstance(value, (list, tuple, set)):
            for item in value:
                devices.update(self._extract_devices(item))
            return devices

        for attr_name in ("device", "pnicDevice", "uplinkPort", "uplinkPortKey", "name"):
            attr_value = getattr(value, attr_name, None)
            if isinstance(attr_value, str) and attr_value.startswith("vmnic"):
                devices.add(attr_value)

        for attr_name in ("pnicSpec", "pnic", "backing"):
            nested = getattr(value, attr_name, None)
            if nested is not value:
                devices.update(self._extract_devices(nested))

        return devices

    def _get_actual_speed(self, pnic):
        speed_obj = getattr(pnic, "linkSpeed", None)
        return self._format_speed(speed_obj)

    def _get_configured_speed(self, pnic):
        for candidate in (getattr(pnic, "spec", None), pnic):
            if not candidate:
                continue
            for attr_name in ("linkSpeed", "speed"):
                speed_obj = getattr(candidate, attr_name, None)
                if speed_obj:
                    formatted = self._format_speed(speed_obj)
                    if formatted != "未知":
                        return formatted

        auto_negotiate = getattr(pnic, "autoNegotiateSupported", None)
        if auto_negotiate is True:
            return "自动协商"
        return "自动协商"

    def _format_speed(self, speed_obj):
        if speed_obj is None:
            return "未知"

        speed_mb = getattr(speed_obj, "speedMb", None)
        if speed_mb is None:
            speed_mb = getattr(speed_obj, "speed", None)
        if speed_mb is None:
            if isinstance(speed_obj, (int, float)):
                speed_mb = speed_obj
            else:
                return "未知"

        try:
            speed_mb = float(speed_mb)
        except Exception:
            return "未知"

        if speed_mb <= 0:
            return "未知"
        if speed_mb >= 1000:
            speed_gb = speed_mb / 1000.0
            if abs(speed_gb - round(speed_gb)) < 1e-6:
                return f"{int(round(speed_gb))} Gbps"
            return f"{speed_gb:.1f} Gbps"
        if abs(speed_mb - round(speed_mb)) < 1e-6:
            return f"{int(round(speed_mb))} Mbps"
        return f"{speed_mb:.1f} Mbps"

    def _get_switch_name(self, device, pnic, switch_map):
        switch_names = sorted(switch_map.get(device, set()))
        if switch_names:
            return "；".join(switch_names)

        for attr_name in ("vswitch", "switch", "portgroup", "dvs", "distributedVirtualSwitch"):
            value = getattr(pnic, attr_name, None)
            if value is None:
                continue
            if isinstance(value, str) and value.strip():
                return value.strip()
            if hasattr(value, "name") and getattr(value, "name", ""):
                return str(getattr(value, "name", "")).strip()

        return "未知"

    def _nic_sort_key(self, device):
        match = re.search(r"(\d+)$", device)
        if match:
            prefix = device[: match.start(1)]
            return (prefix, int(match.group(1)))
        return (device, 10**9)

    def save_to_excel(self, file_path, rows):
        region_map = self.db.get_vcenter_region_map()
        wb = Workbook()
        ws = wb.active
        ws.title = "ESXi网卡巡检"

        headers = ["区域", "vCenter", "主机名", "集群", "设备", "实际速度", "已配置的速度", "交换机"]
        ws.append(headers)

        header_fill = PatternFill(fill_type="solid", fgColor="4472C4")
        header_font = Font(bold=True, color="FFFFFF")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

        grouped_rows = defaultdict(list)
        for row in rows:
            group_key = (
                str(row.get("vc_ip", "")),
                str(row.get("host_name", "")),
                str(row.get("cluster", "")),
                str(region_map.get(str(row.get("vc_ip", "")), "")),
            )
            grouped_rows[group_key].append(row)

        merge_ranges = []
        current_row = 2
        for group_key in grouped_rows:
            vc_ip, host_name, cluster_name, region = group_key
            group_rows = grouped_rows[group_key]
            group_size = len(group_rows)

            for idx, row in enumerate(group_rows):
                excel_row = current_row + idx
                ws.cell(row=excel_row, column=1, value=region)
                ws.cell(row=excel_row, column=2, value=vc_ip)
                ws.cell(row=excel_row, column=3, value=host_name)
                ws.cell(row=excel_row, column=4, value=cluster_name)
                ws.cell(row=excel_row, column=5, value=row.get("device", ""))
                ws.cell(row=excel_row, column=6, value=row.get("actual_speed", ""))
                ws.cell(row=excel_row, column=7, value=row.get("configured_speed", ""))
                ws.cell(row=excel_row, column=8, value=row.get("switch", ""))

                for col in range(1, 9):
                    ws.cell(row=excel_row, column=col).alignment = center

            if group_size > 1:
                start_row = current_row
                end_row = current_row + group_size - 1
                for col in range(1, 5):
                    merge_ranges.append((start_row, col, end_row, col))

            current_row += group_size

        for start_row, start_col, end_row, end_col in merge_ranges:
            ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
            ws.cell(row=start_row, column=start_col).alignment = center

        widths = {
            "A": 12,
            "B": 15,
            "C": 24,
            "D": 22,
            "E": 14,
            "F": 16,
            "G": 18,
            "H": 26,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        wb.save(file_path)
        format_excel_file(file_path, max_width=36)
        return file_path
