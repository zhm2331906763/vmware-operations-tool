import ssl
import concurrent.futures
from pyVim.connect import SmartConnect, Disconnect
from pyVmomi import vim
from openpyxl import Workbook
from openpyxl.styles import Alignment


class ISCSIExporter:
    def __init__(self):
        self.results = []
    
    def export_iscsi(self, vcenters, progress_callback=None, status_callback=None):
        self.results = []
        
        def process_single_vc(vc):
            vc_ip, user, pwd = vc[1], vc[2], vc[3]
            context = ssl._create_unverified_context()
            si = None
            vc_results = []
            
            try:
                if status_callback:
                    status_callback(f"正在处理 vCenter: {vc_ip}")
                
                si = SmartConnect(host=vc_ip, user=user, pwd=pwd, port=443, sslContext=context)
                content = si.content
                
                container = content.viewManager.CreateContainerView(
                    content.rootFolder, [vim.HostSystem], True
                )
                
                hosts = list(container.view)
                container.Destroy()
                
                for host in hosts:
                    try:
                        host_name = host.name
                        cluster_name = self._get_cluster_name(host)
                        host_iqn, send_target, static_target = self._get_iscsi_info(host)
                        
                        vc_results.append({
                            'vCenter': vc_ip,
                            '集群': cluster_name,
                            '主机': host_name,
                            'IQN标识符': host_iqn,
                            '动态目标(SendTarget)': send_target,
                            '静态目标(StaticTarget)': static_target
                        })
                    except Exception:
                        continue
                
                if si:
                    Disconnect(si)
                    si = None
                
                return vc_results
            
            except Exception as e:
                if si:
                    Disconnect(si)
                return vc_results
        
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
            futures = {executor.submit(process_single_vc, vc): vc for vc in vcenters}
            for future in concurrent.futures.as_completed(futures):
                results = future.result()
                self.results.extend(results)
        
        return self.results
    
    def _get_cluster_name(self, host):
        try:
            parent = host.parent
            while parent:
                if isinstance(parent, vim.ClusterComputeResource):
                    return parent.name
                parent = parent.parent
            return '无独立集群'
        except:
            return '未知'
    
    def _get_iscsi_info(self, host):
        try:
            if not hasattr(host, 'configManager') or not host.configManager:
                return '未找到配置管理器', '', ''
            
            storage_system = host.configManager.storageSystem
            if not storage_system:
                return '未找到存储系统', '', ''
            
            device_info = storage_system.storageDeviceInfo
            if not device_info or not device_info.hostBusAdapter:
                return '未找到存储适配器', '', ''
            
            host_iqns = []
            send_targets = []
            static_targets = []
            
            for hba in device_info.hostBusAdapter:
                if isinstance(hba, vim.host.InternetScsiHba):
                    if hasattr(hba, 'iScsiName') and hba.iScsiName:
                        host_iqns.append(hba.iScsiName)
                    
                    if hasattr(hba, 'configuredSendTarget') and hba.configuredSendTarget:
                        for target in hba.configuredSendTarget:
                            if hasattr(target, 'address') and target.address:
                                send_targets.append(target.address)
                    
                    if hasattr(hba, 'configuredStaticTarget') and hba.configuredStaticTarget:
                        for target in hba.configuredStaticTarget:
                            if hasattr(target, 'iScsiName') and target.iScsiName:
                                static_targets.append(target.iScsiName)
            
            host_iqn_str = ', '.join(host_iqns) if host_iqns else '未配置iSCSI'
            send_target_str = '\n'.join(sorted(set(send_targets))) if send_targets else '未配置'
            static_target_str = '\n'.join(sorted(set(static_targets))) if static_targets else '未配置'
            
            return host_iqn_str, send_target_str, static_target_str
        
        except Exception as e:
            return f'获取失败: {str(e)}', '', ''
    
    def save_to_excel(self, file_path):
        if not self.results:
            return False
        
        wb = Workbook()
        ws = wb.active
        ws.title = "ESXi_iSCSI信息"
        
        headers = ['vCenter', '集群', '主机', 'IQN标识符', '动态目标(SendTarget)', '静态目标(StaticTarget)']
        ws.append(headers)
        
        wrap_alignment = Alignment(wrap_text=True, vertical='top')
        
        for result in self.results:
            ws.append([
                result['vCenter'],
                result['集群'],
                result['主机'],
                result['IQN标识符'],
                result['动态目标(SendTarget)'],
                result['静态目标(StaticTarget)']
            ])
            
            last_row = ws.max_row
            ws.cell(row=last_row, column=5).alignment = wrap_alignment
            ws.cell(row=last_row, column=6).alignment = wrap_alignment
        
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 60
        ws.column_dimensions['E'].width = 45
        ws.column_dimensions['F'].width = 60
        
        wb.save(file_path)
        return True