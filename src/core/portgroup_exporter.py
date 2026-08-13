import ssl
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pyVim.connect import Disconnect, SmartConnect
from pyVmomi import vim, vmodl

from src.core.excel_formatter import format_excel_file
from src.core.nsx_utils import is_nsx_portgroup_name
from src.db.database import Database


class PortGroupExporter:
    def _connect_vc(self, vc_ip, username, password):
        context = ssl._create_unverified_context()
        return SmartConnect(
            host=vc_ip,
            user=username,
            pwd=password,
            port=443,
            sslContext=context,
        )

    def _create_property_filter_spec(self, container, vim_type, properties):
        try:
            obj_spec = vmodl.query.PropertyCollector.ObjectSpec()
            obj_spec.obj = container
            obj_spec.selectSet = [
                vmodl.query.PropertyCollector.TraversalSpec(
                    name="traverseEntities",
                    path="view",
                    skip=False,
                    type=vim.view.ContainerView,
                )
            ]

            prop_spec = vmodl.query.PropertyCollector.PropertySpec()
            prop_spec.type = vim_type
            prop_spec.pathSet = properties

            filter_spec = vmodl.query.PropertyCollector.FilterSpec()
            filter_spec.objectSet = [obj_spec]
            filter_spec.propSet = [prop_spec]
            return filter_spec
        except Exception:
            return None

    def _batch_get_datacenters(self, content):
        dc_map = {}
        try:
            container = content.viewManager.CreateContainerView(content.rootFolder, [vim.Datacenter], True)
            for dc in container.view:
                dc_map[dc._moId] = dc.name
            container.Destroy()
        except Exception:
            pass
        return dc_map

    def _batch_get_clusters(self, content):
        cluster_map = {}
        try:
            container = content.viewManager.CreateContainerView(
                content.rootFolder, [vim.ClusterComputeResource], True
            )
            for cluster in container.view:
                cluster_map[cluster._moId] = cluster.name
            container.Destroy()
        except Exception:
            pass
        return cluster_map

    def _batch_get_vds(self, content):
        vds_map = {}
        try:
            container = content.viewManager.CreateContainerView(
                content.rootFolder, [vim.dvs.VmwareDistributedVirtualSwitch], True
            )
            for vds in container.view:
                vds_map[vds._moId] = {"name": vds.name, "cluster": ""}
            container.Destroy()
        except Exception:
            pass
        return vds_map

    def _get_vlan_id(self, config):
        try:
            if config and hasattr(config, "defaultPortConfig"):
                port_config = config.defaultPortConfig
                if hasattr(port_config, "vlan"):
                    vlan_spec = port_config.vlan
                    if hasattr(vlan_spec, "vlanId"):
                        vlan_value = vlan_spec.vlanId
                        if isinstance(vlan_value, int):
                            return vlan_value
                        if isinstance(vlan_value, (list, tuple)) and len(vlan_value) > 0:
                            return "0-4094"
                        if hasattr(vlan_value, "start") and hasattr(vlan_value, "end"):
                            return "0-4094"
                        if hasattr(vlan_value, "__class__") and "NumericRange" in str(vlan_value.__class__):
                            return "0-4094"
                        return str(vlan_value) if vlan_value else 0
            return 0
        except Exception:
            return 0

    def _is_nsx_portgroup(self, pg_name, config=None, vds_name=""):
        if is_nsx_portgroup_name(pg_name) or is_nsx_portgroup_name(vds_name):
            return True
        name = (pg_name or "").strip().lower()
        vds = (vds_name or "").strip().lower()
        if name.startswith("vxw-") or vds.startswith("vxw-"):
            return True
        if config and hasattr(config, "backingType"):
            backing_type = str(getattr(config, "backingType", "") or "").lower()
            if "nsx" in backing_type or "opaque" in backing_type:
                return True
        return False

    def collect_portgroups(self, vc_ip, username, password):
        results = []
        si = None

        try:
            si = self._connect_vc(vc_ip, username, password)
            content = si.content

            self._cluster_cache = self._batch_get_clusters(content)
            vds_map = self._batch_get_vds(content)

            container = content.viewManager.CreateContainerView(
                content.rootFolder, [vim.dvs.DistributedVirtualPortgroup], True
            )

            properties = ["name", "config", "parent"]
            filter_spec = self._create_property_filter_spec(container, vim.dvs.DistributedVirtualPortgroup, properties)
            pc = content.propertyCollector
            result = pc.RetrieveProperties(specSet=[filter_spec])
            container.Destroy()

            for obj_content in result:
                try:
                    props = {p.name: p.val for p in obj_content.propSet}
                    pg_name = props.get("name", "")

                    vds_name = ""
                    cluster_name = ""
                    config = props.get("config")
                    if config and hasattr(config, "distributedVirtualSwitch"):
                        vds = config.distributedVirtualSwitch
                        if vds and vds._moId in vds_map:
                            vds_info = vds_map[vds._moId]
                            vds_name = vds_info["name"]
                            cluster_name = vds_info["cluster"]

                    parent_obj = props.get("parent")
                    datacenter_name = ""
                    if parent_obj:
                        try:
                            while parent_obj:
                                if isinstance(parent_obj, vim.Datacenter):
                                    datacenter_name = parent_obj.name
                                    break
                                parent_obj = parent_obj.parent if hasattr(parent_obj, "parent") else None
                        except Exception:
                            pass

                    vlan_id = self._get_vlan_id(config)
                    is_nsx = self._is_nsx_portgroup(pg_name, config, vds_name)

                    results.append(
                        {
                            "vCenter": vc_ip,
                            "数据中心": datacenter_name,
                            "集群": cluster_name,
                            "端口组名称": pg_name,
                            "端口组VLAN": vlan_id,
                            "所属交换机": vds_name,
                            "端口组类型": "NSX端口组" if is_nsx else "普通端口组",
                            "is_nsx": is_nsx,
                        }
                    )
                except Exception:
                    continue

            if si:
                Disconnect(si)
                si = None

        except Exception:
            if si:
                try:
                    Disconnect(si)
                except Exception:
                    pass

        return results

    def export_all_vcenters(self, vcenters, status_callback=None):
        all_results = []

        def collect_single(vc_config):
            vc_ip = vc_config[1]
            username = vc_config[2]
            password = vc_config[3]
            start_time = time.time()
            results = self.collect_portgroups(vc_ip, username, password)
            elapsed = time.time() - start_time
            return results, elapsed

        if status_callback:
            status_callback(f"开始收集 {len(vcenters)} 个vCenter...")

        max_workers = min(5, len(vcenters)) or 1
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {executor.submit(collect_single, vc): vc for vc in vcenters}
            total_time = 0
            for future in as_completed(futures):
                try:
                    results, elapsed = future.result()
                    all_results.extend(results)
                    total_time += elapsed
                    if status_callback:
                        status_callback(f"已收集 {len(all_results)} 个端口组 (耗时: {elapsed:.1f}s)")
                except Exception:
                    continue

        return all_results, total_time

    def save_to_excel(self, file_path, results):
        region_map = Database().get_vcenter_region_map()
        wb = Workbook()
        ws = wb.active
        ws.title = "端口组信息"

        headers = ["区域", "vCenter", "数据中心", "集群", "端口组名称", "端口组VLAN", "所属交换机", "端口组类型"]
        header_fill = PatternFill(fill_type="solid", fgColor="4472C4")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=column, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        nsx_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")

        for r in results:
            current_row = ws.max_row + 1
            values = [
                region_map.get(str(r.get("vCenter", "")), ""),
                r.get("vCenter", ""),
                r.get("数据中心", ""),
                r.get("集群", ""),
                r.get("端口组名称", ""),
                r.get("端口组VLAN", ""),
                r.get("所属交换机", ""),
                r.get("端口组类型", ""),
            ]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row=current_row, column=column, value=value)
                cell.alignment = row_alignment

            if r.get("is_nsx"):
                for col in range(1, len(headers) + 1):
                    ws.cell(row=current_row, column=col).fill = nsx_fill

        column_widths = {"A": 12, "B": 15, "C": 20, "D": 20, "E": 30, "F": 15, "G": 25, "H": 16}
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        ws.row_dimensions[1].height = 30
        wb.save(file_path)
        format_excel_file(file_path)
