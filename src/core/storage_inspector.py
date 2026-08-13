import concurrent.futures

from pyVmomi import vim
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill, Font

from src.core.concurrent_config import ConcurrentConfig
from src.core.excel_formatter import format_excel_file


class StorageInspector:
    REQUIRED_COLUMNS = [
        '区域', 'Vcenter', '集群', '存储名称', '存储类型', '文件系统类型',
        '文件系统版本', '总容量(GB)', '可用容量(GB)', '可用容量百分比(%)',
        '是否为SSD', '块大小/MB', '整体状态', '事件告警数',
        '设备详细名称', '设备UUID', '数据存储路径'
    ]

    def __init__(self, connector):
        self.connector = connector

    def collect_all_datastores(self, skip_inaccessible=False):
        if not self.connector.is_connected():
            return []

        ds_properties = [
            'name',
            'summary',
            'info',
            'host',
            'overallStatus',
            'triggeredAlarmState',
        ]
        ds_result = self.connector._retrieve_properties_all(vim.Datastore, ds_properties)
        if not ds_result:
            return []

        rows = []
        workers = ConcurrentConfig.get_optimal_workers('host_collect_parse', len(ds_result))
        workers = max(4, min(workers, 32))

        def parse_one(obj_content):
            datastore = obj_content.obj
            props = {p.name: p.val for p in obj_content.propSet}
            if skip_inaccessible and not self._is_datastore_accessible(datastore, props):
                return None

            summary = props.get('summary')
            ds_name = props.get('name', '')
            total_bytes = summary.capacity if summary else 0
            free_bytes = summary.freeSpace if summary else 0

            ds_info = {
                '区域': '',
                'Vcenter': self.connector.ip,
                '集群': self._get_cluster_name_from_datastore(datastore, props),
                '存储名称': ds_name,
                '存储类型': self._get_storage_type_fast(datastore, props),
                '文件系统类型': self._get_fs_type(datastore, props),
                '文件系统版本': self._get_fs_version(datastore, props),
                '总容量(GB)': round(total_bytes / (1024 ** 3), 2) if total_bytes > 0 else 0,
                '可用容量(GB)': round(free_bytes / (1024 ** 3), 2) if free_bytes > 0 else 0,
                '可用容量百分比(%)': round((free_bytes / total_bytes) * 100, 1) if total_bytes > 0 else 0,
                '是否为SSD': self._is_ssd_datastore(datastore, props),
                '块大小/MB': self._get_block_size(datastore, props),
                '整体状态': self._get_overall_status(datastore, props),
                '事件告警数': self._get_alarm_count(props),
                '设备详细名称': self._get_device_name(datastore, props),
                '设备UUID': self._get_device_uuid(datastore, props),
                '数据存储路径': self._get_datastore_path(datastore, props),
            }

            fs_type_upper = str(ds_info.get('文件系统类型', '')).upper()
            if fs_type_upper in ('NFS', 'NFS41'):
                ds_info['是否为SSD'] = '否'
                ds_info['设备详细名称'] = ''
                ds_info['设备UUID'] = ''

            if fs_type_upper == 'VSAN':
                ds_info['设备详细名称'] = ''
                ds_info['设备UUID'] = ''

            return ds_info

        with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
            futures = [executor.submit(parse_one, obj_content) for obj_content in ds_result]
            for future in concurrent.futures.as_completed(futures):
                try:
                    row = future.result()
                    if row:
                        rows.append(row)
                except Exception:
                    continue

        rows.sort(key=lambda x: (x.get('Vcenter', ''), x.get('集群', ''), x.get('存储名称', '')))
        return rows

    @staticmethod
    def _get_parent_name(obj, target_type):
        try:
            parent = obj.parent
            while parent:
                if isinstance(parent, target_type):
                    return parent.name
                if hasattr(parent, 'parent'):
                    parent = parent.parent
                else:
                    break
            return ''
        except Exception:
            return ''

    def _get_storage_type_fast(self, datastore, props):
        try:
            summary = props.get('summary')
            info = props.get('info')
            ds_type = ''
            if summary and hasattr(summary, 'type'):
                ds_type = str(summary.type)
            ds_type_norm = ds_type.lower()

            if info and hasattr(info, 'nas'):
                return '共享'

            if ds_type_norm in ['nfs', 'nfs41', 'vsan', 'vvol']:
                return '共享'

            if summary and hasattr(summary, 'multipleHostAccess'):
                if summary.multipleHostAccess is False:
                    return '本地'
                if summary.multipleHostAccess is True:
                    return '共享'

            if ds_type_norm == 'vmfs':
                host_count = len(datastore.host) if hasattr(datastore, 'host') and datastore.host else 0
                if host_count > 1:
                    return '共享'
                if host_count == 1:
                    return '本地'
                if info and hasattr(info, 'vmfs') and getattr(info.vmfs, 'extent', None):
                    return '本地'
                return '未知'

            return '未知'
        except Exception:
            return '未知'

    @staticmethod
    def _is_ssd_datastore(datastore, props):
        try:
            info = props.get('info')
            if info and hasattr(info, 'vmfs') and hasattr(info.vmfs, 'ssd'):
                return '是' if info.vmfs.ssd else '否'

            ds_name = (props.get('name') or '').lower()
            if any(kw in ds_name for kw in {'ssd', 'nvme', 'flash'}):
                return '是(推测)'

            if info and hasattr(info, 'nas'):
                return 'NFS无法判断'

            return '否'
        except Exception:
            return '未知'

    @staticmethod
    def _get_block_size(datastore, props):
        try:
            info = props.get('info')
            if info and hasattr(info, 'vmfs') and hasattr(info.vmfs, 'blockSizeMb'):
                return info.vmfs.blockSizeMb
            return 'N/A'
        except Exception:
            return 'N/A'

    @staticmethod
    def _get_fs_version(datastore, props):
        try:
            info = props.get('info')
            summary = props.get('summary')

            if info and hasattr(info, 'vmfs') and hasattr(info.vmfs, 'version'):
                return info.vmfs.version
            if info and hasattr(info, 'nas'):
                return 'NFS'
            if summary and hasattr(summary, 'type') and str(summary.type).lower() == 'vsan':
                return 'VSAN'
            if summary and hasattr(summary, 'type') and str(summary.type).lower() == 'vvol':
                return 'vVOL'
            return 'N/A'
        except Exception:
            return 'N/A'

    def _get_fs_type(self, datastore, props):
        try:
            summary = props.get('summary')
            if not summary or not hasattr(summary, 'type'):
                return 'N/A'
            ds_type = str(summary.type)
            if ds_type.lower() == 'vsan':
                return 'vsan'
            return ds_type
        except Exception:
            return 'N/A'

    def _get_cluster_name_from_datastore(self, datastore, props):
        try:
            if not hasattr(datastore, 'host') or not datastore.host:
                return ''

            clusters = []
            for host_mount in datastore.host:
                host = getattr(host_mount, 'key', None)
                if not host:
                    continue
                cluster_name = self._get_parent_name(host, vim.ClusterComputeResource)
                if cluster_name and cluster_name not in clusters:
                    clusters.append(cluster_name)
            return '\n'.join(clusters)
        except Exception:
            return ''

    @staticmethod
    def _get_device_uuid(datastore, props):
        try:
            info = props.get('info')
            summary = props.get('summary')
            if info and hasattr(info, 'vmfs') and hasattr(info.vmfs, 'uuid'):
                return info.vmfs.uuid
            if summary and hasattr(summary, 'url'):
                return summary.url
            return 'N/A'
        except Exception:
            return 'N/A'

    @staticmethod
    def _get_datastore_path(datastore, props):
        try:
            summary = props.get('summary')
            if summary and hasattr(summary, 'url'):
                return summary.url
            return 'N/A'
        except Exception:
            return 'N/A'

    @staticmethod
    def _get_device_name(datastore, props):
        try:
            info = props.get('info')
            if info and hasattr(info, 'vmfs') and hasattr(info.vmfs, 'extent') and info.vmfs.extent:
                extent = info.vmfs.extent[0]
                if hasattr(extent, 'diskName'):
                    return extent.diskName
            if info and hasattr(info, 'nas'):
                if hasattr(info.nas, 'remoteHost') and hasattr(info.nas, 'remotePath'):
                    return f"{info.nas.remoteHost}:{info.nas.remotePath}"
            return props.get('name', '')
        except Exception:
            return props.get('name', '')

    @staticmethod
    def _get_overall_status(datastore, props):
        try:
            status_val = props.get('overallStatus', getattr(datastore, 'overallStatus', ''))
            status_map = {
                'green': '正常',
                'yellow': '警告',
                'red': '严重',
                'gray': '未知',
            }
            resolved = status_map.get(str(status_val), str(status_val) if status_val else '')
            if resolved and resolved != '未知':
                return resolved

            summary = props.get('summary')
            if summary and hasattr(summary, 'accessible') and not summary.accessible:
                return '严重'

            alarm_state = props.get('triggeredAlarmState', [])
            if alarm_state and len(alarm_state) > 0:
                return '警告'

            if summary and hasattr(summary, 'accessible') and summary.accessible:
                return '正常'

            return '未知'
        except Exception:
            return '未知'

    @staticmethod
    def _is_datastore_accessible(datastore, props):
        try:
            summary = props.get('summary', getattr(datastore, 'summary', None))
            if summary is not None and hasattr(summary, 'accessible'):
                return bool(summary.accessible)
            return True
        except Exception:
            return True

    @staticmethod
    def _get_alarm_count(props):
        alarm_state = props.get('triggeredAlarmState', [])
        return len(alarm_state) if alarm_state else 0

    @classmethod
    def save_to_excel(cls, file_path, all_ds_info, region_map=None):
        wb = Workbook()
        ws = wb.active
        ws.title = '存储信息'

        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        red_font = Font(color='9C0006')
        yellow_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        yellow_font = Font(color='9C5700')

        headers = cls.REQUIRED_COLUMNS
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_num, ds_info in enumerate(all_ds_info or [], 2):
            for col_num, header in enumerate(headers, 1):
                value = ds_info.get(header, '')
                if header == '区域' and not value and region_map:
                    value = region_map.get(str(ds_info.get('Vcenter', '')), '')
                ws.cell(row=row_num, column=col_num, value=value)
                ws.cell(row=row_num, column=col_num).alignment = center_align

            percent_cell = ws.cell(row=row_num, column=10)
            try:
                available_percent = float(str(percent_cell.value).strip().replace('%', '') or 0)
            except Exception:
                available_percent = None

            if str(ws.cell(row=row_num, column=5).value).strip() == '共享' and available_percent is not None and available_percent < 5:
                percent_cell.fill = red_fill
                percent_cell.font = red_font

            status_cell = ws.cell(row=row_num, column=13)
            status = str(status_cell.value)
            if status == '严重':
                status_cell.fill = red_fill
                status_cell.font = red_font
            elif status == '警告':
                status_cell.fill = yellow_fill
                status_cell.font = yellow_font

        column_widths = {
            'A': 15, 'B': 20, 'C': 30, 'D': 12, 'E': 15, 'F': 15,
            'G': 15, 'H': 15, 'I': 20, 'J': 12, 'K': 12, 'L': 12,
            'M': 12, 'N': 50, 'O': 40, 'P': 50, 'Q': 50
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'A2'

        # Storage inspection is intentionally a single-sheet report. Remove any
        # unexpected worksheets before and after formatting so an old workbook
        # or formatter cannot bring back the retired CVM sheet.
        for extra_sheet in list(wb.worksheets):
            if extra_sheet is not ws:
                wb.remove(extra_sheet)
        wb.active = 0
        wb.save(file_path)
        format_excel_file(file_path, sheet_name='存储信息')

        formatted_wb = load_workbook(file_path)
        for extra_sheet in list(formatted_wb.worksheets):
            if extra_sheet.title != '存储信息':
                formatted_wb.remove(extra_sheet)
        formatted_wb.active = 0
        formatted_wb.save(file_path)
        return True
