from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_HEADER_FILL = '4472C4'
DEFAULT_HEADER_FONT = 'FFFFFF'
DEFAULT_HIGHLIGHT_FILL = 'E2F0D9'
DEFAULT_ROW_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
DEFAULT_HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
DEFAULT_THIN_SIDE = Side(style='thin', color='D9D9D9')
DEFAULT_BORDER = Border(
    left=DEFAULT_THIN_SIDE,
    right=DEFAULT_THIN_SIDE,
    top=DEFAULT_THIN_SIDE,
    bottom=DEFAULT_THIN_SIDE,
)


def _display_length(value):
    if value is None:
        return 0
    text = str(value)
    if not text:
        return 0
    return max(len(line) for line in text.splitlines())


def format_excel_file(
    file_path,
    sheet_name=None,
    highlight_column=None,
    highlight_values=None,
    highlight_fill=DEFAULT_HIGHLIGHT_FILL,
    max_width=48,
    min_width=10,
):
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    headers = [cell.value for cell in ws[1]]
    header_fill = PatternFill(fill_type='solid', fgColor=DEFAULT_HEADER_FILL)
    header_font = Font(bold=True, color=DEFAULT_HEADER_FONT)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = DEFAULT_HEADER_ALIGN
        cell.border = DEFAULT_BORDER

    highlight_idx = None
    normalized_highlights = None
    if highlight_column and highlight_column in headers:
        highlight_idx = headers.index(highlight_column) + 1
    if highlight_values:
        normalized_highlights = {str(v).strip().lower() for v in highlight_values}

    widths = {}
    for idx, header in enumerate(headers, 1):
        widths[idx] = max(min_width, _display_length(header) + 2)

    green_fill = PatternFill(fill_type='solid', fgColor=highlight_fill)

    for row_num, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        row_height = 18
        highlight_row = False

        if highlight_idx and normalized_highlights is not None:
            value = row[highlight_idx - 1].value if len(row) >= highlight_idx else None
            if str(value).strip().lower() in normalized_highlights:
                highlight_row = True

        for idx, cell in enumerate(row, 1):
            if isinstance(cell, MergedCell):
                continue
            if cell.value is None:
                cell.value = ''
            cell.alignment = DEFAULT_ROW_ALIGN
            cell.border = DEFAULT_BORDER
            widths[idx] = max(widths.get(idx, min_width), min(max_width, _display_length(cell.value) + 2))
            if isinstance(cell.value, str):
                row_height = max(row_height, min(72, 18 * max(1, cell.value.count('\n') + 1)))

        if highlight_row:
            for cell in row:
                cell.fill = green_fill

        ws.row_dimensions[row_num].height = row_height

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(file_path)
    return file_path


def highlight_column_cells(
    file_path,
    column_name,
    target_values,
    sheet_name=None,
    fill_color='FFC7CE',
    font_color='9C0006',
):
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]
    if column_name not in headers:
        wb.save(file_path)
        return file_path

    col_idx = headers.index(column_name) + 1
    normalized_targets = {str(v).strip().lower() for v in target_values}
    cell_fill = PatternFill(fill_type='solid', fgColor=fill_color)
    cell_font = Font(color=font_color)

    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if str(cell.value).strip().lower() in normalized_targets:
            cell.fill = cell_fill
            cell.font = cell_font

    wb.save(file_path)
    return file_path
